"""Checks a generated workbook against the excel-generation data rules.

Read-only. Reports problems that make a delivered workbook unreliable or
awkward to use, grouped by sheet with cell references. Exits 1 if any ERROR is
found, so it can gate delivery.

ERROR (fix before delivering):

- `numbers-as-text`: a mostly numeric column (years, amounts, counts,
  percentages, dates) stores values as text, which breaks sorting, filtering,
  and SUM.
- `leaked-token`: a cell contains a code artifact such as "None", "nan", or
  "<empty>".
- `formula-error`: a cell's cached value is an Excel error such as #REF!.
- `header`: a data sheet has a blank or duplicate header.
- `too-long`: a cell exceeds Excel's 32,767-character limit.

WARN (fix unless there is a reason not to):

- `id-as-number`: an identifier column holds long or decimal numbers, so digits
  or leading zeros may be lost; identifiers belong in text cells.
- `mixed-types`: a column mixes numbers or dates with text.
- `missing-markers`: a column mixes blanks with markers like "N/A" or "—".
- `whitespace`, `sheet-name`, `no-freeze`, `autofilter`, `merged-cells`,
  `column-widths`, `fonts`: layout and consistency problems.

INFO:

- `uncached-formulas`: formulas have no cached values yet, so previews and
  pandas show blanks until the file is recalculated (see recalc.py) or opened
  in Excel.

Typical usage example:

  python check_workbook.py Acme_Vehicle_Reconciliation_2026-09-11.xlsx
"""

import argparse
import datetime as dt
import json
import re
import sys
from collections import Counter
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Final, Literal, cast

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

type Severity = Literal["ERROR", "WARN", "INFO"]

_EXCEL_ERRORS: Final = frozenset(
    {
        "#REF!",
        "#DIV/0!",
        "#NAME?",
        "#VALUE!",
        "#N/A",
        "#NUM!",
        "#NULL!",
        "#SPILL!",
        "#CALC!",
    }
)
_LEAKED_TOKENS: Final = frozenset(
    {"none", "nan", "nat", "null", "undefined", "<empty>", "<na>", "[object object]"}
)
# Placeholders that leak into longer strings, e.g. "2005 PJ <empty>".
_EMBEDDED_TOKEN: Final = re.compile(
    r"<(empty|na|null|none|nan|missing)>|\[object Object\]", re.IGNORECASE
)
_MISSING_MARKERS: Final = frozenset(
    {"n/a", "na", "-", "--", "\u2013", "\u2014", "tbd", "unknown", "?"}
)
_NUMERIC_TEXT: Final = (
    re.compile(r"^\(?-?\$?\s*-?\d{1,3}(,\d{3})+(\.\d+)?\)?$"),  # 1,250 / $4,000
    re.compile(r"^\(?-?\$?\s*-?\d+(\.\d+)?\)?$"),  # 2004 / $400 / (12.5)
    re.compile(r"^-?\d+(\.\d+)?\s*%$"),  # 94%
    re.compile(r"^\d{1,2}/\d{1,2}/\d{2,4}$"),  # 02/09/2026
    re.compile(r"^\d{4}-\d{2}-\d{2}$"),  # 2026-02-09
)
# Headers naming an identifier rather than a quantity: numeric-looking text is
# correct there ("Unit #" 007, "Model" 2500, "ZIP" 02134).
_ID_HEADER: Final = re.compile(
    r"#|(?<![a-z])(vin|serial|s/n|id|no\.?|num|number|code|zip|postal|phone|fax|"
    r"policy|claim|tag|unit|account|acct|fein|ein|naics|sic|model|license|plate|"
    r"ref|reference|row|page)(?![a-z])"
)
# Words that make a header a quantity even if it also names an entity
# ("Claim Amount", "Policy Year", "Unit Cost").
_MEASURE_HEADER: Final = re.compile(
    r"amount|amt|premium|cost|value|price|limit|deductible|count|total|paid|"
    r"reserve|incurred|rate|ratio|percent|pct|%|year|date|age|tiv|payroll|"
    r"revenue|sales|miles|sq|area|score|similarity"
)
_NUMERIC_SHARE_FOR_ERROR: Final = 0.8
_EXCEL_CELL_LIMIT: Final = 32_767
_EXCEL_GENERAL_DIGITS: Final = 11
_MAX_EXAMPLES: Final = 3
# Written by xlsx_kit.add_matrix_sheet: comparison matrices legitimately mix
# numbers and text per column and have no autofilter.
_MATRIX_SHEETS_PROPERTY: Final = "excel-generation:matrix-sheets"


@dataclass(frozen=True, slots=True)
class Issue:
    """One problem found in the workbook.

    Attributes:
        severity: ERROR, WARN, or INFO.
        sheet: Sheet title, or an empty string for workbook-wide issues.
        rule: Short rule identifier, e.g. `numbers-as-text`.
        message: What is wrong, where, and how to fix it.
    """

    severity: Severity
    sheet: str
    rule: str
    message: str


def check_workbook(path: Path) -> list[Issue]:
    """Runs every check on a workbook.

    Args:
        path: The .xlsx file to check.

    Returns:
        Issues ordered by sheet, then severity.
    """
    formulas_wb = load_workbook(path)
    values_wb = load_workbook(path, data_only=True)
    issues: list[Issue] = []
    fonts: Counter[str] = Counter()
    matrix_sheets = _matrix_sheets(formulas_wb)

    for ws in formulas_wb.worksheets:
        cached = values_wb[ws.title]
        is_matrix = ws.title in matrix_sheets
        issues.extend(_check_sheet_name(ws.title))
        issues.extend(_check_cells(ws, cached, fonts))
        if _is_table_sheet(ws):
            issues.extend(_check_table_layout(ws, is_matrix=is_matrix))
            issues.extend(_check_columns(ws, is_matrix=is_matrix))
        else:
            issues.extend(_check_loose_numbers(ws))

    if len(fonts) > 1:
        listed = ", ".join(f"{name} ({count})" for name, count in fonts.most_common())
        message = f"Cells use {len(fonts)} font families: {listed}. Use one throughout."
        issues.append(Issue("WARN", "", "fonts", message))
    order = {"ERROR": 0, "WARN": 1, "INFO": 2}
    return sorted(issues, key=lambda issue: (issue.sheet, order[issue.severity]))


def _matrix_sheets(wb: Workbook) -> set[str]:
    """Returns titles of sheets that xlsx_kit marked as comparison matrices.

    Args:
        wb: The loaded workbook.
    """
    # openpyxl's type stubs predate custom document properties.
    props = cast(Any, wb).custom_doc_props
    for prop in props.props:
        if prop.name == _MATRIX_SHEETS_PROPERTY:
            try:
                return set(json.loads(str(prop.value)))
            except json.JSONDecodeError:
                return set()
    return set()


def _check_sheet_name(title: str) -> list[Issue]:
    """Flags sheet names that break references or import tools."""
    if title.isascii() and len(title) <= 31:
        return []
    message = (
        f"Sheet name {title!r} contains emoji or other non-ASCII characters, or "
        "exceeds 31 characters; use plain ASCII so formulas and imports work."
    )
    return [Issue("WARN", title, "sheet-name", message)]


def _check_cells(ws: Worksheet, cached: Worksheet, fonts: Counter[str]) -> list[Issue]:
    """Checks every non-empty cell for leaks, errors, length, and caching.

    Args:
        ws: The sheet loaded with formulas.
        cached: The same sheet loaded with cached values.
        fonts: Running count of font families, updated in place.

    Returns:
        Cell-level issues for this sheet.
    """
    leaked: list[str] = []
    errors: list[str] = []
    too_long: list[str] = []
    uncached = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or isinstance(cell, MergedCell):
                continue
            # openpyxl's stubs leave Font.name untyped; it is a str or None.
            font_name = cast(str | None, cell.font.name)  # pyright: ignore[reportUnknownMemberType]
            fonts[font_name or "default"] += 1
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                cached_value = cached[cell.coordinate].value
                if cached_value is None:
                    uncached += 1
                elif isinstance(cached_value, str) and cached_value in _EXCEL_ERRORS:
                    errors.append(f"{cell.coordinate}={cached_value}")
                continue
            if isinstance(value, str):
                if value.strip().lower() in _LEAKED_TOKENS or _EMBEDDED_TOKEN.search(
                    value
                ):
                    leaked.append(f"{cell.coordinate}={value!r}")
                elif value.strip() in _EXCEL_ERRORS:
                    errors.append(f"{cell.coordinate}={value}")
                if len(value) > _EXCEL_CELL_LIMIT:
                    too_long.append(cell.coordinate)

    issues: list[Issue] = []
    if leaked:
        message = (
            f"{len(leaked)} cell(s) contain code artifacts ({_examples(leaked)}). "
            "Write a blank cell instead."
        )
        issues.append(Issue("ERROR", ws.title, "leaked-token", message))
    if errors:
        message = f"{len(errors)} cell(s) show Excel errors ({_examples(errors)})."
        issues.append(Issue("ERROR", ws.title, "formula-error", message))
    if too_long:
        message = f"Cells over {_EXCEL_CELL_LIMIT:,} characters: {_examples(too_long)}."
        issues.append(Issue("ERROR", ws.title, "too-long", message))
    if uncached:
        message = (
            f"{uncached} formula(s) have no cached value; previews will show blanks "
            "until the file is recalculated (recalc.py) or opened in Excel."
        )
        issues.append(Issue("INFO", ws.title, "uncached-formulas", message))
    return issues


def _is_table_sheet(ws: Worksheet) -> bool:
    """Returns whether the sheet is a data table with its header in row 1."""
    header = [cell.value for cell in ws[1] if cell.value is not None]
    return len(header) >= 2 and ws.max_row >= 2


def _check_table_layout(ws: Worksheet, *, is_matrix: bool) -> list[Issue]:
    """Checks header, freeze panes, autofilter, merges, and column widths.

    Args:
        ws: A sheet whose header is in row 1.
        is_matrix: Whether the sheet is a comparison matrix, which needs no
            autofilter.

    Returns:
        Layout issues for this sheet.
    """
    issues: list[Issue] = []
    headers = [cell.value for cell in ws[1]][: ws.max_column]
    blank = [
        get_column_letter(i) for i, value in enumerate(headers, 1) if value is None
    ]
    duplicated = sorted(
        str(value) for value, count in Counter(headers).items() if value and count > 1
    )
    if blank or duplicated:
        details: list[str] = []
        if blank:
            details.append(f"blank headers in column(s) {', '.join(blank)}")
        if duplicated:
            details.append(f"duplicate headers {', '.join(duplicated)}")
        issues.append(Issue("ERROR", ws.title, "header", "; ".join(details) + "."))

    if ws.freeze_panes is None:
        message = "Header row isn't frozen; set freeze panes at A2."
        issues.append(Issue("WARN", ws.title, "no-freeze", message))

    filter_ref = ws.auto_filter.ref
    if is_matrix:
        pass
    elif not filter_ref:
        message = "No autofilter; apply one over the full data range."
        issues.append(Issue("WARN", ws.title, "autofilter", message))
    else:
        _, _, max_col, max_row = range_boundaries(filter_ref)
        if (max_col or 0) < ws.max_column or (max_row or 0) < ws.max_row:
            full = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
            message = f"Autofilter covers {filter_ref} but the data spans {full}."
            issues.append(Issue("WARN", ws.title, "autofilter", message))

    if ws.merged_cells.ranges:
        merged = [str(cell_range) for cell_range in ws.merged_cells]
        message = (
            f"Merged cells in a data sheet ({_examples(merged)}) break sorting, "
            "filtering, and pasting; unmerge them."
        )
        issues.append(Issue("WARN", ws.title, "merged-cells", message))

    if not any(dim.width for dim in ws.column_dimensions.values()):
        message = "Column widths are all default; size columns to their content."
        issues.append(Issue("WARN", ws.title, "column-widths", message))
    return issues


def _check_columns(ws: Worksheet, *, is_matrix: bool) -> list[Issue]:
    """Profiles each column's values and flags type and consistency problems.

    Args:
        ws: A sheet whose header is in row 1.
        is_matrix: Whether the sheet is a comparison matrix, where columns
            legitimately mix numbers with text such as "Excluded".

    Returns:
        Column issues for this sheet.
    """
    issues: list[Issue] = []
    for index, header_cell in enumerate(ws[1], start=1):
        header = str(header_cell.value or get_column_letter(index))
        cells = [
            cell
            for (cell,) in ws.iter_rows(min_row=2, min_col=index, max_col=index)
            if isinstance(cell, Cell)
        ]
        issues.extend(_check_column(ws.title, header, cells, is_matrix=is_matrix))
    return issues


def _check_column(
    sheet: str, header: str, cells: Sequence[Cell], *, is_matrix: bool = False
) -> list[Issue]:
    """Checks one data column.

    Args:
        sheet: Sheet title, used to label issues.
        header: The column's header text.
        cells: Data cells below the header.
        is_matrix: Skip the mixed-type and missing-marker checks, which don't
            apply to comparison matrices.

    Returns:
        Issues for this column.
    """
    values = [
        cell
        for cell in cells
        if cell.value is not None
        and not (isinstance(cell.value, str) and cell.value.startswith("="))
    ]
    if not values:
        return []
    header_key = header.lower()
    is_id = bool(_ID_HEADER.search(header_key)) and not _MEASURE_HEADER.search(
        header_key
    )
    texts = [cell for cell in values if isinstance(cell.value, str)]
    numbers = [cell for cell in values if isinstance(cell.value, int | float | dt.date)]
    numeric_texts = [cell for cell in texts if _looks_numeric(str(cell.value))]
    markers = [
        cell for cell in texts if str(cell.value).strip().lower() in _MISSING_MARKERS
    ]
    label = f"{header!r}"
    issues: list[Issue] = []

    if is_id:
        risky = [
            cell
            for cell in values
            if isinstance(cell.value, float)
            or (
                isinstance(cell.value, int)
                and len(str(abs(cell.value))) > _EXCEL_GENERAL_DIGITS
            )
        ]
        if risky:
            message = (
                f"{label} is an identifier but {len(risky)} cell(s) hold long or "
                f"decimal numbers ({_cell_examples(risky)}); store identifiers as text."
            )
            issues.append(Issue("WARN", sheet, "id-as-number", message))
    elif numeric_texts:
        # "No value" markers are judged separately, so they don't dilute the share.
        measured = len(values) - len(markers)
        numeric_share = (len(numeric_texts) + len(numbers)) / measured
        if numeric_share >= _NUMERIC_SHARE_FOR_ERROR:
            message = (
                f"{len(numeric_texts)} cell(s) in {label} store numbers as text "
                f"({_cell_examples(numeric_texts)}); write numbers with a number "
                "format."
            )
            issues.append(Issue("ERROR", sheet, "numbers-as-text", message))

    non_numeric_texts = [cell for cell in texts if cell not in numeric_texts]
    if is_matrix:
        pass
    elif numbers and non_numeric_texts and not is_id:
        stray = [cell for cell in non_numeric_texts if cell not in markers] or markers
        message = (
            f"{label} mixes numbers with text ({_cell_examples(stray)}); leave unknown "
            "values blank and move explanations to a Notes column."
        )
        issues.append(Issue("WARN", sheet, "mixed-types", message))
    if markers and len(values) < len(cells) and not is_matrix:
        message = (
            f"{label} mixes blank cells with 'no value' markers "
            f"({_cell_examples(markers)}); use one convention, preferably blanks."
        )
        issues.append(Issue("WARN", sheet, "missing-markers", message))

    padded = [cell for cell in texts if str(cell.value) != str(cell.value).strip()]
    if padded:
        message = f"{len(padded)} cell(s) in {label} have leading or trailing spaces."
        issues.append(Issue("WARN", sheet, "whitespace", message))
    return issues


def _check_loose_numbers(ws: Worksheet) -> list[Issue]:
    """Flags numeric-looking text on non-table sheets such as the summary.

    A label to the left that names an identifier ("Policy #") exempts the cell.
    """
    flagged: list[str] = []
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell, Cell) or not isinstance(cell.value, str):
                continue
            text = cell.value.strip()
            if not _looks_numeric(text):
                continue
            label = (
                ws.cell(row=cell.row, column=cell.column - 1).value
                if cell.column > 1
                else None
            )
            if isinstance(label, str) and _ID_HEADER.search(label.lower()):
                continue
            flagged.append(f"{cell.coordinate}={text!r}")
    if not flagged:
        return []
    message = (
        f"{len(flagged)} cell(s) hold numbers as text ({_examples(flagged)}); "
        "write numbers (or formulas) with a number format."
    )
    return [Issue("WARN", ws.title, "numbers-as-text", message)]


def _looks_numeric(text: str) -> bool:
    """Returns whether text is a number, amount, percentage, or date.

    Integers with a leading zero or more than 15 digits are treated as
    identifiers, not numbers.
    """
    stripped = text.strip()
    digits = stripped.replace(",", "")
    if digits.isdigit() and (
        len(digits) > 15 or (len(digits) > 1 and digits.startswith("0"))
    ):
        return False
    return any(pattern.match(stripped) for pattern in _NUMERIC_TEXT)


def _cell_examples(cells: Sequence[Cell]) -> str:
    """Formats a few cells as `A2='value'` examples."""
    return _examples([f"{cell.coordinate}={cell.value!r}" for cell in cells])


def _examples(items: Sequence[str]) -> str:
    """Joins the first few items, noting how many more there are."""
    shown = ", ".join(items[:_MAX_EXAMPLES])
    remaining = len(items) - _MAX_EXAMPLES
    return f"{shown}, +{remaining} more" if remaining > 0 else shown


def _parse_args(argv: Sequence[str] | None) -> argparse.Namespace:
    """Parses command-line arguments."""
    parser = argparse.ArgumentParser(description="Check a generated workbook.")
    parser.add_argument("workbook", type=Path, help="The .xlsx file to check.")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    """Checks a workbook and prints the issues found.

    Args:
        argv: Command-line arguments, excluding the program name. Defaults to
            `sys.argv[1:]`.

    Returns:
        Process exit code: 1 if any ERROR was found, otherwise 0.
    """
    args = _parse_args(argv)
    issues = check_workbook(args.workbook)
    for issue in issues:
        where = issue.sheet or "workbook"
        print(f"[{issue.severity}] {where} ({issue.rule}): {issue.message}")
    counts = Counter(issue.severity for issue in issues)
    print(
        f"{args.workbook.name}: {counts['ERROR']} error(s), "
        f"{counts['WARN']} warning(s), {counts['INFO']} info"
    )
    return 1 if counts["ERROR"] else 0


if __name__ == "__main__":
    sys.exit(main())
