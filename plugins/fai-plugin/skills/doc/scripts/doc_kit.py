#!/usr/bin/env python3
"""doc_kit.py — FurtherAI .docx layout library.

Encodes the brand DOCX grammar measured from the two reference templates in
``assets/brand/doc_templates/``:

* **client** grammar (``Pricing_Document.docx``) — the brand-correct,
  client-facing look: square logo, Fraunces Light green title, Wix Madefor
  Display body, and a borderless table with a green header band.
* **internal** grammar (``Meeting_Brief.docx``) — an unbranded Arial prep
  sheet: bold blue title, italic meta line, blue underlined headings, a
  two-column fact table and bullet lists.

Every helper writes direct run formatting rather than relying on Word styles,
because neither template's styles carry the brand — the templates format
everything inline, and so do we.

Requires ``python-docx`` (``python3 -m pip install python-docx``). Pillow is optional and
only used to crop the wordmark logo.
"""
from __future__ import annotations

import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Iterable, Optional, Sequence

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:  # pragma: no cover - dependency guard
    sys.stderr.write(
        "doc_kit requires python-docx.\n"
        "  python3 -m pip install python-docx\n"
    )
    raise SystemExit(1)

PLUGIN_ROOT = Path(__file__).resolve().parents[3]
BRAND = PLUGIN_ROOT / "assets" / "brand"
FONT_DIR = BRAND / "fonts"
LOGO_DIR = BRAND / "logos"

# --------------------------------------------------------------------------
# Fonts — shared resolver lives in the deck skill; fall back if absent.
# --------------------------------------------------------------------------


class _FallbackFonts:
    """Minimal stand-in for ``skills/deck/scripts/fonts.py``.

    TODO: delete once the deck skill's shared resolver is on disk; this only
    exists so the doc skill works standalone. Same ``FAMILY`` / ``check()`` /
    ``install()`` surface.
    """

    FAMILY = {
        "display_light": "Fraunces 72pt Light",
        "display": "Fraunces 72pt",
        "display_semibold": "Fraunces 72pt SemiBold",
        "body": "Wix Madefor Display",
        "body_medium": "Wix Madefor Display Medium",
        "body_semibold": "Wix Madefor Display SemiBold",
        "body_extrabold": "Wix Madefor Display ExtraBold",
        "mono": "Courier New",
    }

    def check(self):
        """Return {family_name: installed?} for the brand families."""
        import subprocess

        try:
            out = subprocess.run(
                ["fc-list", ":", "family"], capture_output=True, text=True, timeout=20
            ).stdout
        except Exception:
            out = ""
        if not out:
            try:  # macOS has no fc-list by default; read the user font dir.
                out = "\n".join(p.name for p in Path.home().joinpath("Library/Fonts").glob("*.ttf"))
            except Exception:
                out = ""
        return {
            name: (name.replace(" ", "") in out.replace(" ", ""))
            for name in sorted(set(self.FAMILY.values()))
        }

    def install(self):
        """Copy the static brand TTFs into the user font directory (macOS)."""
        import shutil

        dest = Path.home() / "Library" / "Fonts"
        if not dest.exists():
            return []
        copied = []
        for ttf in _all_static_ttfs():
            target = dest / ttf.name
            if not target.exists():
                shutil.copy2(ttf, target)
                copied.append(target)
        return copied


def _all_static_ttfs():
    yield from (FONT_DIR / "Fraunces" / "static").glob("Fraunces_72pt-*.ttf")
    yield from (FONT_DIR / "Wix_Madefor_Display" / "static").glob("*.ttf")


def _load_fonts():
    deck_scripts = PLUGIN_ROOT / "skills" / "deck" / "scripts"
    if (deck_scripts / "fonts.py").exists():
        sys.path.insert(0, str(deck_scripts))
        try:
            import fonts as _fonts  # type: ignore

            if hasattr(_fonts, "FAMILY"):
                return _fonts
        except Exception:  # pragma: no cover - never let a broken import kill a build
            pass
    return _FallbackFonts()


fonts = _load_fonts()


def family(key: str) -> str:
    """Resolve a font-role key through the shared resolver, tolerating a
    resolver whose FAMILY map is keyed differently."""
    fam = getattr(fonts, "FAMILY", {}) or {}
    if key in fam:
        return fam[key]
    return _FallbackFonts.FAMILY[key]


FONTS = {k: family(k) for k in _FallbackFonts.FAMILY}

#: Font families used by the client grammar.
DISPLAY_LIGHT = FONTS["display_light"]
DISPLAY = FONTS["display"]
BODY = FONTS["body"]
BODY_SEMIBOLD = FONTS["body_semibold"]
#: The internal grammar is deliberately unbranded.
INTERNAL_FONT = "Arial"

# --------------------------------------------------------------------------
# Tokens — Palette 1 ("Product / brand"), §A.4 of the template spec.
# --------------------------------------------------------------------------

PALETTE = {
    "green_deep": "#074B40",
    "green": "#25654F",
    "green_dark_text": "#1D4438",
    "green_pale": "#D9EAD3",
    "ink": "#14161C",
    "ink_alt": "#2B2D31",
    "warm_dark": "#444339",
    "paper": "#FBFBF9",
    "paper_warm": "#F4F3F0",
    "card": "#F8F7F5",
    "tan": "#D0CEC3",
    "hair": "#CCCCCC",
    "muted": "#595959",
    "white": "#FFFFFF",
    "black": "#000000",
    # Internal (Meeting_Brief) grammar.
    "internal_blue": "#1F4E79",
    "internal_label": "#EEF3F8",
    "internal_muted": "#555555",
    "internal_rule": "#CCCCCC",
}


def rgb(color: str) -> RGBColor:
    """'#074B40' or '074B40' -> RGBColor."""
    return RGBColor.from_string(color.lstrip("#").upper())


#: The palette as RGBColor values, for callers that want them pre-built.
RGB = {k: rgb(v) for k, v in PALETTE.items()}

#: Margin presets measured from the two templates: (left, right, top, bottom).
MARGINS = {
    "client": (1.0, 1.0, 0.9, 0.9),
    "internal": (0.75, 0.75, 0.75, 0.75),
}

#: Content width in inches for each preset (8.5in page minus L/R margins).
CONTENT_WIDTH = {"client": 6.5, "internal": 7.0}

_TWIPS_PER_INCH = 1440
#: Pricing_Document cell margins, in twips: top / left / bottom / right.
CELL_MARGINS = (140, 200, 140, 120)
#: Meeting_Brief fact-table cell margins, in twips.
FACT_CELL_MARGINS = (60, 120, 60, 120)

# --------------------------------------------------------------------------
# Low-level OXML helpers
# --------------------------------------------------------------------------


def _el(tag: str, **attrs) -> "OxmlElement":
    node = OxmlElement(tag)
    for name, value in attrs.items():
        node.set(qn("w:" + name), str(value))
    return node


def _hex(color: str) -> str:
    if color in ("auto", "none"):
        return color
    return color.lstrip("#").upper()


def set_font(run, name: str, size: Optional[float] = None, bold: bool = False,
             italic: bool = False, color: Optional[str] = None):
    """Apply a font to a run, writing ``w:rFonts`` on the run XML directly.

    ``run.font.name`` alone sets only ``w:ascii`` on some python-docx paths and
    does not survive a Word round-trip for non-ASCII-mapped families, so every
    helper in this module goes through here: ascii / hAnsi / cs / eastAsia are
    all set to the same family name.
    """
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = rgb(color)
    return run


def _spacing(paragraph, before: Optional[float] = None, after: Optional[float] = None,
             line: Optional[float] = None):
    pf = paragraph.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    return paragraph


def _bottom_border(paragraph, color: str, sz: int = 6, space: int = 1):
    """Add a paragraph bottom rule (the Meeting_Brief heading underline)."""
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        # pBdr must precede spacing/ind in the CT_PPr sequence.
        anchor = ppr.find(qn("w:spacing"))
        if anchor is not None:
            anchor.addprevious(pbdr)
        else:
            ppr.append(pbdr)
    bottom = _el("w:bottom", val="single", sz=sz, space=space, color=_hex(color))
    pbdr.append(bottom)
    return paragraph


def _shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = _el("w:shd", val="clear", color="auto", fill=_hex(fill))
    tc_pr.append(shd)


def _cell_margins(cell, margins=CELL_MARGINS):
    top, left, bottom, right = margins
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement("w:" + tag)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def _cell_width(cell, width_tw: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcw = tc_pr.get_or_add_tcW()
    tcw.set(qn("w:w"), str(int(width_tw)))
    tcw.set(qn("w:type"), "dxa")


def _table_borders(outer: Optional[dict] = None, inside: Optional[dict] = None):
    """Build a ``w:tblBorders`` element. Each spec dict is {'val','sz','color'}."""
    borders = OxmlElement("w:tblBorders")
    outer = outer or {"val": "nil", "sz": 0, "color": "000000"}
    inside = inside or {"val": "single", "sz": 4, "color": "000000"}
    for edge in ("top", "left", "bottom", "right"):
        borders.append(_el("w:" + edge, val=outer["val"], sz=outer["sz"],
                           space=0, color=_hex(outer["color"])))
    for edge in ("insideH", "insideV"):
        borders.append(_el("w:" + edge, val=inside["val"], sz=inside["sz"],
                           space=0, color=_hex(inside["color"])))
    return borders


def _cell_borders(cell, spec: dict):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        borders.append(_el("w:" + edge, val=spec["val"], sz=spec["sz"],
                           space=0, color=_hex(spec["color"])))
    tc_pr.append(borders)


def _grid_twips(col_widths_in: Sequence[float], target_total_in: Optional[float] = None):
    """Column widths in twips, apportioned so they sum to the target width.

    Largest-remainder apportionment: the spec's inch values are themselves
    rounded from twips, so naive rounding drifts a twip or two off the 6.5in
    content width. This lands 1.701/1.701/3.097in on exactly 2450/2450/4460.
    """
    raw = [w * _TWIPS_PER_INCH for w in col_widths_in]
    total = sum(raw)
    if target_total_in is not None and abs(sum(col_widths_in) - target_total_in) <= 0.05:
        total = target_total_in * _TWIPS_PER_INCH
        scale = total / sum(raw)
        raw = [r * scale for r in raw]
    target = int(round(total))
    tw = [int(r) for r in raw]
    order = sorted(range(len(raw)), key=lambda i: raw[i] - tw[i], reverse=True)
    for i in order[: target - sum(tw)]:
        tw[i] += 1
    return tw


def _configure_table(table, col_widths_in: Sequence[float],
                     outer: Optional[dict] = None, inside: Optional[dict] = None,
                     target_total_in: Optional[float] = None):
    """Rebuild ``w:tblPr`` and ``w:tblGrid`` for a fixed-layout table.

    Children are emitted in CT_TblPr schema order (tblW, tblBorders, tblLayout,
    tblLook) so Word accepts the file without repair.
    """
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    for child in list(tbl_pr):
        tbl_pr.remove(child)
    widths = _grid_twips(col_widths_in, target_total_in)
    tbl_pr.append(_el("w:tblW", w=sum(widths), type="dxa"))
    tbl_pr.append(_table_borders(outer, inside))
    tbl_pr.append(_el("w:tblLayout", type="fixed"))
    tbl_pr.append(_el("w:tblLook", val="0000"))

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl_pr.addnext(grid)
    table.autofit = False
    return widths


def _header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(_el("w:tblHeader", val="1"))


def _cell_text(cell, text: str, font: str, size: float, color: str,
               bold: bool = False, italic: bool = False, space_after: float = 0):
    """Fill a cell, splitting on newlines into separate paragraphs."""
    lines = str(text).split("\n")
    first = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for run in list(first.runs):
        run._r.getparent().remove(run._r)
    paragraphs = [first]
    for _ in lines[1:]:
        paragraphs.append(cell.add_paragraph())
    for paragraph, line in zip(paragraphs, lines):
        _spacing(paragraph, after=space_after)
        run = paragraph.add_run(line)
        set_font(run, font, size=size, bold=bold, italic=italic, color=color)
    return cell


# --------------------------------------------------------------------------
# Document setup
# --------------------------------------------------------------------------


def _set_doc_defaults(doc, name: str, size: float, color: str):
    styles = doc.styles.element
    docdefaults = styles.find(qn("w:docDefaults"))
    if docdefaults is None:
        docdefaults = OxmlElement("w:docDefaults")
        styles.insert(0, docdefaults)
    rpr_default = docdefaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        docdefaults.insert(0, rpr_default)
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    for child in list(rpr):
        rpr.remove(child)
    rfonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    rpr.append(rfonts)
    rpr.append(_el("w:color", val=_hex(color)))
    rpr.append(_el("w:sz", val=int(round(size * 2))))
    rpr.append(_el("w:szCs", val=int(round(size * 2))))
    # Keep the Normal style in step so python-docx-created paragraphs inherit.
    normal = doc.styles["Normal"]
    normal.font.name = name
    normal.font.size = Pt(size)
    normal.font.color.rgb = rgb(color)
    n_rpr = normal.element.get_or_add_rPr()
    n_fonts = n_rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        n_fonts.set(qn(attr), name)


def new_doc(margins="client", default_font: Optional[str] = None,
            default_size: float = 11, default_color: str = "#2B2D31"):
    """A blank US-Letter portrait document with brand docDefaults.

    ``margins`` is a preset name (``'client'`` = L/R 1.0in T/B 0.9in, matching
    the client-facing template; ``'internal'`` = 0.75in all round) or an
    explicit ``(left, right, top, bottom)`` tuple of inches.
    ``default_font`` defaults to Wix Madefor Display for the client preset and
    Arial for the internal one.
    """
    if isinstance(margins, str):
        preset = margins
        left, right, top, bottom = MARGINS[margins]
    else:
        preset = "custom"
        left, right, top, bottom = margins
    if default_font is None:
        default_font = INTERNAL_FONT if preset == "internal" else BODY
    if preset == "internal" and default_color == "#2B2D31":
        default_color = "#000000"
        default_size = 10 if default_size == 11 else default_size

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    _set_doc_defaults(doc, default_font, default_size, default_color)
    doc._fai = {
        "preset": preset,
        "font": default_font,
        "size": default_size,
        "color": default_color,
        "content_width": CONTENT_WIDTH.get(preset, 8.5 - left - right),
    }
    return doc


def _meta(doc, key, fallback):
    return getattr(doc, "_fai", {}).get(key, fallback)


# --------------------------------------------------------------------------
# Logo
# --------------------------------------------------------------------------

#: Content bounding box of Logo_Full_Green_on_White.jpg — the shipped file is
#: ~57% padding, so the wordmark must be cropped before placement.
WORDMARK_BBOX = (22, 116, 614, 242)
#: Default placed width for the square mark and the horizontal wordmark.
SQUARE_WIDTH_IN = 0.646
WORDMARK_WIDTH_IN = 1.6
_CROP_CACHE = Path(tempfile.gettempdir()) / "fai_doc_kit"


def _wordmark_path():
    """Crop the wordmark JPEG to its content bbox. Returns (path, cropped?)."""
    src = LOGO_DIR / "Logo_Full_Green_on_White.jpg"
    try:
        from PIL import Image
    except ImportError:
        return None, False
    _CROP_CACHE.mkdir(parents=True, exist_ok=True)
    out = _CROP_CACHE / "wordmark_cropped.png"
    if not out.exists():
        with Image.open(src) as img:
            img.convert("RGB").crop(WORDMARK_BBOX).save(out, "PNG")
    return out, True


def logo(doc, variant: str = "square", width_in: float = SQUARE_WIDTH_IN,
         space_after: float = 12):
    """Place the brand logo as an inline picture in its own paragraph.

    ``variant='square'`` uses the square green mark (the client template's
    0.646 x 0.646in slot). ``variant='wordmark'`` uses the horizontal
    mark+wordmark, cropped to its content bounding box with Pillow; without
    Pillow it falls back to the square mark and says so on stderr.

    ``width_in`` defaults to the square slot; a wordmark left at that default
    is placed at ``WORDMARK_WIDTH_IN`` instead, since 0.646in would be
    unreadably small for a 4.7:1 mark.
    """
    paragraph = doc.add_paragraph()
    _spacing(paragraph, after=space_after)
    run = paragraph.add_run()
    if variant == "wordmark":
        path, cropped = _wordmark_path()
        if path is None:
            sys.stderr.write(
                "doc_kit: Pillow not installed, cannot crop the wordmark "
                "(the shipped JPEG is ~57% padding) — using the square mark "
                "instead. Run `python3 -m pip install Pillow` for the wordmark.\n"
            )
            path, cropped = LOGO_DIR / "Logo_Social_Green_on_White.png", False
        if cropped:
            aspect = (WORDMARK_BBOX[2] - WORDMARK_BBOX[0]) / (WORDMARK_BBOX[3] - WORDMARK_BBOX[1])
            width = WORDMARK_WIDTH_IN if width_in == SQUARE_WIDTH_IN else width_in
            run.add_picture(str(path), width=Inches(width), height=Inches(width / aspect))
        else:
            run.add_picture(str(path), width=Inches(width_in), height=Inches(width_in))
    else:
        path = LOGO_DIR / "Logo_Social_Green_on_White.png"
        run.add_picture(str(path), width=Inches(width_in), height=Inches(width_in))
    return paragraph


# --------------------------------------------------------------------------
# Client grammar (Pricing_Document.docx)
# --------------------------------------------------------------------------


def title(doc, text: str, space_after: float = 4):
    """Fraunces Light 32pt deep green — the document title."""
    paragraph = doc.add_paragraph()
    _spacing(paragraph, after=space_after)
    set_font(paragraph.add_run(text), DISPLAY_LIGHT, size=32, color=PALETTE["green_deep"])
    return paragraph


def subtitle(doc, text: str, space_after: float = 3):
    """Wix Madefor Display 11pt gray — the 'Prepared for … · date' line."""
    paragraph = doc.add_paragraph()
    _spacing(paragraph, after=space_after)
    set_font(paragraph.add_run(text), BODY, size=11, color=PALETTE["muted"])
    return paragraph


def body(doc, text: str, space_after: float = 8, size: float = 11,
         color: str = PALETTE["ink_alt"], bold: bool = False, italic: bool = False):
    """Wix Madefor Display 11pt charcoal body paragraph."""
    paragraph = doc.add_paragraph()
    _spacing(paragraph, after=space_after)
    set_font(paragraph.add_run(text), _meta(doc, "font", BODY), size=size,
             bold=bold, italic=italic, color=color)
    return paragraph


def section_heading(doc, text: str, space_before: float = 6, space_after: float = 7):
    """Fraunces Light 22pt deep green — the section heading above a table."""
    paragraph = doc.add_paragraph()
    _spacing(paragraph, before=space_before, after=space_after)
    set_font(paragraph.add_run(text), DISPLAY_LIGHT, size=22, color=PALETTE["green_deep"])
    return paragraph


def rule(doc, color: str = PALETTE["tan"], space_after: float = 14, sz: int = 6):
    """An empty paragraph carrying a bottom rule — the template's divider."""
    paragraph = doc.add_paragraph()
    _spacing(paragraph, after=space_after)
    _bottom_border(paragraph, color, sz=sz, space=0)
    return paragraph


def brand_table(doc, headers: Sequence[str], rows: Sequence[Sequence[str]],
                col_widths_in: Sequence[float], total_row: bool = False):
    """The reusable brand table: no outer frame, hairline interior rules, a
    deep-green header band, an off-white first column and an optional pale
    green total row.

    Column typography follows the template: column 1 is Wix Madefor Display
    SemiBold 10.5pt charcoal on ``#FBFBF9``; column 3 is 10pt gray; every other
    body column is Wix Madefor Display 10.5pt charcoal. Cell strings may
    contain ``\\n`` to produce multiple paragraphs inside one cell.
    """
    n_cols = len(headers)
    if any(len(r) != n_cols for r in rows):
        raise ValueError(
            f"brand_table: every row must have {n_cols} cells to match the headers"
        )
    if len(col_widths_in) != n_cols:
        raise ValueError(
            f"brand_table: col_widths_in must have {n_cols} entries, got {len(col_widths_in)}"
        )

    table = doc.add_table(rows=len(rows) + 1, cols=n_cols)
    widths = _configure_table(
        table, col_widths_in,
        outer={"val": "nil", "sz": 0, "color": "000000"},
        inside={"val": "single", "sz": 4, "color": "000000"},
        target_total_in=_meta(doc, "content_width", None),
    )

    header = table.rows[0]
    _header_row(header)
    for idx, cell in enumerate(header.cells):
        _cell_width(cell, widths[idx])
        _shade(cell, PALETTE["green_deep"])
        _cell_margins(cell, CELL_MARGINS)
        _cell_text(cell, headers[idx], BODY_SEMIBOLD, 10.5, PALETTE["white"], bold=True)

    last = len(rows) - 1
    for r_idx, row_values in enumerate(rows):
        is_total = total_row and r_idx == last
        for c_idx, cell in enumerate(table.rows[r_idx + 1].cells):
            _cell_width(cell, widths[c_idx])
            _cell_margins(cell, CELL_MARGINS)
            if is_total:
                _shade(cell, PALETTE["green_pale"])
                font_name, size, color = BODY_SEMIBOLD, 10.5, PALETTE["green_deep"]
            elif c_idx == 0:
                _shade(cell, PALETTE["paper"])
                font_name, size, color = BODY_SEMIBOLD, 10.5, PALETTE["ink_alt"]
            elif c_idx == 2:
                font_name, size, color = BODY, 10, PALETTE["muted"]
            else:
                font_name, size, color = BODY, 10.5, PALETTE["ink_alt"]
            _cell_text(cell, row_values[c_idx], font_name, size, color)
    return table


# --------------------------------------------------------------------------
# Internal grammar (Meeting_Brief.docx)
# --------------------------------------------------------------------------


def internal_title(doc, text: str, space_after: float = 2):
    """Arial 15pt bold `#1F4E79` — the internal brief's title line."""
    paragraph = doc.add_paragraph()
    _spacing(paragraph, after=space_after)
    set_font(paragraph.add_run(text), INTERNAL_FONT, size=15, bold=True,
             color=PALETTE["internal_blue"])
    return paragraph


def internal_heading(doc, text: str, space_before: float = 10, space_after: float = 5):
    """Arial 12pt bold `#1F4E79` with the blue bottom rule under it."""
    paragraph = doc.add_paragraph()
    _spacing(paragraph, before=space_before, after=space_after)
    _bottom_border(paragraph, PALETTE["internal_blue"], sz=6, space=1)
    set_font(paragraph.add_run(text), INTERNAL_FONT, size=12, bold=True,
             color=PALETTE["internal_blue"])
    return paragraph


def meta_line(doc, text: str, space_after: float = 8):
    """Arial 9pt italic gray — the date / channel / purpose line."""
    paragraph = doc.add_paragraph()
    _spacing(paragraph, after=space_after)
    set_font(paragraph.add_run(text), INTERNAL_FONT, size=9, italic=True,
             color=PALETTE["internal_muted"])
    return paragraph


def sources_line(doc, text: str, space_before: float = 10):
    """Arial 8pt gray closing line; the part before the first ':' is bolded."""
    paragraph = doc.add_paragraph()
    _spacing(paragraph, before=space_before)
    label, sep, rest = str(text).partition(":")
    if sep:
        set_font(paragraph.add_run(label + sep), INTERNAL_FONT, size=8, bold=True,
                 color=PALETTE["internal_muted"])
        set_font(paragraph.add_run(rest), INTERNAL_FONT, size=8,
                 color=PALETTE["internal_muted"])
    else:
        set_font(paragraph.add_run(text), INTERNAL_FONT, size=8,
                 color=PALETTE["internal_muted"])
    return paragraph


def bullets(doc, items: Iterable[str], size: float = 10, space_after: float = 3,
            font: Optional[str] = None, color: Optional[str] = None):
    """A bullet list in the document's body font."""
    font = font or _meta(doc, "font", INTERNAL_FONT)
    color = color or _meta(doc, "color", "#000000")
    out = []
    for item in items:
        try:
            paragraph = doc.add_paragraph(style="List Bullet")
            text = str(item)
        except KeyError:  # template without the built-in list style
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            text = "•  " + str(item)
        _spacing(paragraph, after=space_after)
        set_font(paragraph.add_run(text), font, size=size, color=color)
        out.append(paragraph)
    return out


def label_value_table(doc, pairs: Sequence[Sequence[str]], label_w: float = 1.667,
                      value_w: float = 4.833, size: float = 9.5):
    """The internal fact table: `#EEF3F8` bold label cells, hairline gray grid."""
    table = doc.add_table(rows=len(pairs), cols=2)
    widths = _configure_table(
        table, (label_w, value_w),
        outer={"val": "single", "sz": 4, "color": "auto"},
        inside={"val": "single", "sz": 4, "color": "auto"},
        target_total_in=_meta(doc, "content_width", None),
    )
    hairline = {"val": "single", "sz": 1, "color": PALETTE["internal_rule"]}
    for r_idx, pair in enumerate(pairs):
        label, value = (list(pair) + ["", ""])[:2]
        row = table.rows[r_idx]
        _cell_width(row.cells[0], widths[0])
        _cell_borders(row.cells[0], hairline)
        _shade(row.cells[0], PALETTE["internal_label"])
        _cell_margins(row.cells[0], FACT_CELL_MARGINS)
        _cell_text(row.cells[0], label, INTERNAL_FONT, size, "#000000", bold=True)

        _cell_width(row.cells[1], widths[1])
        _cell_borders(row.cells[1], hairline)
        _cell_margins(row.cells[1], FACT_CELL_MARGINS)
        _cell_text(row.cells[1], value, INTERNAL_FONT, size, "#000000")
    return table


def page_break(doc):
    from docx.enum.text import WD_BREAK

    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    return paragraph


# --------------------------------------------------------------------------
# Font embedding
# --------------------------------------------------------------------------

_FRAUNCES = FONT_DIR / "Fraunces" / "static"
_WMD = FONT_DIR / "Wix_Madefor_Display" / "static"

#: family name -> {embed style: TTF path}. Styles map to the w:embed* elements.
FONT_FILES = {
    "Wix Madefor Display": {
        "Regular": _WMD / "WixMadeforDisplay-Regular.ttf",
        "Bold": _WMD / "WixMadeforDisplay-Bold.ttf",
    },
    "Wix Madefor Display Medium": {
        "Regular": _WMD / "WixMadeforDisplay-Medium.ttf",
        "Bold": _WMD / "WixMadeforDisplay-Bold.ttf",
    },
    "Wix Madefor Display SemiBold": {
        "Regular": _WMD / "WixMadeforDisplay-SemiBold.ttf",
        "Bold": _WMD / "WixMadeforDisplay-Bold.ttf",
    },
    "Wix Madefor Display ExtraBold": {
        "Regular": _WMD / "WixMadeforDisplay-ExtraBold.ttf",
        "Bold": _WMD / "WixMadeforDisplay-ExtraBold.ttf",
    },
    "Fraunces 72pt": {
        "Regular": _FRAUNCES / "Fraunces_72pt-Regular.ttf",
        "Bold": _FRAUNCES / "Fraunces_72pt-Bold.ttf",
        "Italic": _FRAUNCES / "Fraunces_72pt-Italic.ttf",
        "BoldItalic": _FRAUNCES / "Fraunces_72pt-BoldItalic.ttf",
    },
    "Fraunces 72pt Light": {
        "Regular": _FRAUNCES / "Fraunces_72pt-Light.ttf",
        # The Light family has no bold cut; SemiBold is its bold companion.
        "Bold": _FRAUNCES / "Fraunces_72pt-SemiBold.ttf",
        "Italic": _FRAUNCES / "Fraunces_72pt-LightItalic.ttf",
        "BoldItalic": _FRAUNCES / "Fraunces_72pt-SemiBoldItalic.ttf",
    },
    "Fraunces 72pt SemiBold": {
        "Regular": _FRAUNCES / "Fraunces_72pt-SemiBold.ttf",
        "Bold": _FRAUNCES / "Fraunces_72pt-Bold.ttf",
        "Italic": _FRAUNCES / "Fraunces_72pt-SemiBoldItalic.ttf",
        "BoldItalic": _FRAUNCES / "Fraunces_72pt-BoldItalic.ttf",
    },
}

_EMBED_TAG = {
    "Regular": "embedRegular",
    "Bold": "embedBold",
    "Italic": "embedItalic",
    "BoldItalic": "embedBoldItalic",
}

_FONTTABLE_NS = (
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
)
_ZERO_KEY = "{00000000-0000-0000-0000-000000000000}"


def used_families(docx_path) -> list:
    """Family names referenced by w:rFonts in the saved document."""
    import re

    with zipfile.ZipFile(docx_path) as zf:
        blob = b""
        for part in ("word/document.xml", "word/styles.xml"):
            try:
                blob += zf.read(part)
            except KeyError:
                pass
    found = set(re.findall(rb'w:ascii="([^"]+)"', blob))
    return sorted(name.decode("utf-8") for name in found)


def embed_fonts(docx_path, families: Optional[Sequence[str]] = None) -> list:
    """Embed the document's brand fonts into the .docx, the way the client
    template does, so it renders correctly on machines without them installed.

    Rewrites the package in place: adds ``word/fonts/*.ttf``, a fresh
    ``word/fontTable.xml`` with ``w:embedRegular`` / ``w:embedBold`` /
    ``w:embedItalic`` / ``w:embedBoldItalic`` entries, the matching
    ``word/_rels/fontTable.xml.rels``, a ``ttf`` content-type default and
    ``w:embedTrueTypeFonts`` in settings.xml.

    Returns the list of embedded family names. Fonts not vendored under
    ``assets/brand/fonts`` (Arial, Courier New, …) are skipped.
    """
    docx_path = Path(docx_path)
    if families is None:
        families = [f for f in used_families(docx_path) if f in FONT_FILES]
    families = [f for f in families if f in FONT_FILES]
    if not families:
        return []

    with zipfile.ZipFile(docx_path) as zf:
        parts = {name: zf.read(name) for name in zf.namelist()}

    fonts_xml = []
    rels_xml = []
    rid = 0
    embedded = []
    for fam in families:
        entries = []
        for style, src in FONT_FILES[fam].items():
            if not Path(src).exists():
                continue
            rid += 1
            part_name = "word/fonts/{}-{}.ttf".format(fam.replace(" ", ""), style.lower())
            parts[part_name] = Path(src).read_bytes()
            entries.append(
                '<w:{tag} w:fontKey="{key}" r:id="rId{rid}" w:subsetted="0"/>'.format(
                    tag=_EMBED_TAG[style], key=_ZERO_KEY, rid=rid
                )
            )
            rels_xml.append(
                '<Relationship Id="rId{rid}" Type="http://schemas.openxmlformats.org/'
                'officeDocument/2006/relationships/font" Target="fonts/{target}"/>'.format(
                    rid=rid, target=part_name.split("/")[-1]
                )
            )
        if entries:
            fonts_xml.append(
                '<w:font w:name="{name}">{body}</w:font>'.format(name=fam, body="".join(entries))
            )
            embedded.append(fam)

    parts["word/fontTable.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        "<w:fonts {ns}>{body}</w:fonts>".format(ns=_FONTTABLE_NS, body="".join(fonts_xml))
    ).encode("utf-8")
    parts["word/_rels/fontTable.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        "{body}</Relationships>".format(body="".join(rels_xml))
    ).encode("utf-8")

    content_types = parts["[Content_Types].xml"].decode("utf-8")
    if 'Extension="ttf"' not in content_types:
        insert_at = content_types.index(">", content_types.index("<Types ")) + 1
        content_types = (
            content_types[:insert_at]
            + '<Default Extension="ttf" ContentType="application/x-font-ttf"/>'
            + content_types[insert_at:]
        )
    if "/word/fontTable.xml" not in content_types:
        insert_at = content_types.rindex("</Types>")
        content_types = (
            content_types[:insert_at]
            + '<Override PartName="/word/fontTable.xml" ContentType="application/vnd.'
              'openxmlformats-officedocument.wordprocessingml.fontTable+xml"/>'
            + content_types[insert_at:]
        )
    parts["[Content_Types].xml"] = content_types.encode("utf-8")

    settings = parts.get("word/settings.xml", b"").decode("utf-8")
    if settings and "embedTrueTypeFonts" not in settings:
        marker = "<w:defaultTabStop"
        tag = '<w:embedTrueTypeFonts w:val="1"/>'
        if marker in settings:
            settings = settings.replace(marker, tag + marker, 1)
        else:
            insert_at = settings.index(">", settings.index("<w:settings")) + 1
            settings = settings[:insert_at] + tag + settings[insert_at:]
        parts["word/settings.xml"] = settings.encode("utf-8")

    tmp = docx_path.with_suffix(".embedding.tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, blob in parts.items():
            zf.writestr(name, blob)
    tmp.replace(docx_path)
    return embedded


def save(doc, path, embed: bool = True):
    """Save the document and (by default) embed its brand fonts."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    if embed:
        embed_fonts(path)
    return path


if __name__ == "__main__":  # pragma: no cover - smoke check
    print("doc_kit — brand DOCX helpers")
    print("  display light :", DISPLAY_LIGHT)
    print("  display       :", DISPLAY)
    print("  body          :", BODY, "/", BODY_SEMIBOLD)
    print("  font resolver :", "deck/scripts/fonts.py"
          if not isinstance(fonts, _FallbackFonts) else "local fallback")
    print("  margins       :", MARGINS)
