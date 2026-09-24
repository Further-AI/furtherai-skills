"""Caches formula results in a workbook by recalculating it with LibreOffice.

openpyxl writes formulas without results. Excel computes them on open, but
previews, pandas, and many viewers read only cached results and show blanks.
This script opens the workbook in headless LibreOffice, recalculates, saves a
copy, and replaces the original only after verifying that the copy has the
same sheets and the same non-formula values. It then reports any formula that
evaluates to an Excel error.

Exit codes: 0 on success, 1 if formulas evaluate to errors or verification
fails, 2 if LibreOffice isn't installed. Verification failures leave the
original untouched; formula errors are reported after saving the recalculated copy.

Typical usage example:

  python recalc.py Acme_Vehicle_Reconciliation_2026-09-11.xlsx
"""

import argparse
import shutil
import subprocess
import sys
import tempfile
from collections.abc import Sequence
from pathlib import Path
from typing import Final

from openpyxl import load_workbook

_EXCEL_ERRORS: Final = frozenset(
    {"#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!"}
)
# LibreOffice's export filter for the Excel 2007-365 format.
_XLSX_FILTER: Final = "xlsx:Calc MS Excel 2007 XML"


def find_soffice() -> str | None:
    """Returns the path of the LibreOffice executable, or None if absent."""
    return shutil.which("soffice") or shutil.which("libreoffice")


def recalculate(path: Path, *, soffice: str, timeout_seconds: int) -> list[str]:
    """Recalculates a workbook in place.

    Args:
        path: Workbook to recalculate.
        soffice: Path of the LibreOffice executable.
        timeout_seconds: How long to wait for LibreOffice.

    Returns:
        Cells whose formulas evaluate to Excel errors, as `Sheet!A1=#REF!`
        strings. Empty if every formula evaluated cleanly.

    Raises:
        RuntimeError: If LibreOffice fails or its output doesn't match the
            original's data. The original file is not modified in that case.
    """
    with tempfile.TemporaryDirectory(prefix="recalc-") as work:
        work_dir = Path(work)
        out_dir = work_dir / "out"
        # A private profile avoids clashing with any other LibreOffice process.
        profile = (work_dir / "profile").as_uri()
        command = [
            soffice,
            f"-env:UserInstallation={profile}",
            "--headless",
            "--calc",
            "--convert-to",
            _XLSX_FILTER,
            "--outdir",
            str(out_dir),
            str(path.resolve()),
        ]
        try:
            subprocess.run(
                command, capture_output=True, timeout=timeout_seconds, check=True
            )
        except subprocess.TimeoutExpired as err:
            raise RuntimeError(
                f"LibreOffice timed out after {timeout_seconds}s"
            ) from err
        except subprocess.CalledProcessError as err:
            stderr = err.stderr.decode(errors="replace").strip()
            raise RuntimeError(f"LibreOffice failed: {stderr}") from err

        converted = out_dir / path.name
        if not converted.exists():
            raise RuntimeError("LibreOffice produced no output file")
        _verify_same_data(path, converted)
        errors = _formula_errors(converted)
        shutil.copyfile(converted, path)
    return errors


def _verify_same_data(original: Path, converted: Path) -> None:
    """Confirms the recalculated copy kept every sheet and literal value.

    Raises:
        RuntimeError: If sheets or non-formula values differ.
    """
    before = load_workbook(original)
    after = load_workbook(converted)
    if before.sheetnames != after.sheetnames:
        raise RuntimeError(f"Sheets changed: {before.sheetnames} -> {after.sheetnames}")
    for ws in before.worksheets:
        copy = after[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value is None or (isinstance(value, str) and value.startswith("=")):
                    continue
                new_value = copy[cell.coordinate].value
                if new_value != value and str(new_value) != str(value):
                    raise RuntimeError(
                        f"Value changed at {ws.title}!{cell.coordinate}: "
                        f"{value!r} -> {new_value!r}"
                    )


def _formula_errors(path: Path) -> list[str]:
    """Lists cells whose cached formula results are Excel errors."""
    formulas = load_workbook(path)
    values = load_workbook(path, data_only=True)
    errors: list[str] = []
    for ws in formulas.worksheets:
        cached = values[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    continue
                result = cached[cell.coordinate].value
                if isinstance(result, str) and result in _EXCEL_ERRORS:
                    errors.append(f"{ws.title}!{cell.coordinate}={result}")
    return errors


def _parse_args(argv: Sequence[str] | None) -> argparse.Namespace:
    """Parses command-line arguments."""
    parser = argparse.ArgumentParser(
        description="Cache formula results via LibreOffice."
    )
    parser.add_argument("workbook", type=Path, help="The .xlsx file to recalculate.")
    parser.add_argument(
        "--timeout",
        type=int,
        default=120,
        help="Seconds to wait for LibreOffice (default: %(default)s).",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    """Recalculates the workbook and reports the outcome.

    Args:
        argv: Command-line arguments, excluding the program name. Defaults to
            `sys.argv[1:]`.

    Returns:
        Process exit code: 0 on success, 1 on formula errors or failure, 2 if
        LibreOffice isn't installed.
    """
    args = _parse_args(argv)
    soffice = find_soffice()
    if soffice is None:
        print(
            "LibreOffice not found; formulas were not recalculated. They will compute "
            "when the file is opened in Excel, but previews may show blanks."
        )
        return 2
    try:
        errors = recalculate(
            args.workbook, soffice=soffice, timeout_seconds=args.timeout
        )
    except RuntimeError as err:
        print(f"Recalculation failed; original left unchanged. {err}")
        return 1
    if errors:
        shown = ", ".join(errors[:10])
        print(f"Recalculated, but {len(errors)} formula(s) evaluate to errors: {shown}")
        return 1
    print(f"Recalculated {args.workbook.name}; all formulas evaluate cleanly.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
