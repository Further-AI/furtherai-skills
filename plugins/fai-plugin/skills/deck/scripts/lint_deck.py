#!/usr/bin/env python3
"""Static lint for a generated FurtherAI .pptx.

WHY THIS EXISTS
    There is no full-deck renderer on these machines. LibreOffice is not
    installed, so `render_preview.py` falls back to macOS `qlmanage`, which
    rasterises **slide 1 only**. Everything past the first slide ships
    unlooked-at, and the failures that hides are exactly the silent ones:
    text overflowing its box, a shape parked off-canvas, a footer colour
    chosen for a light slide reused on a dark one.

    This lint reads the .pptx structurally and reports what a render would
    have shown. It is not a substitute for looking at a deck — it cannot see
    rhythm, crowding, or whether a layout is any good. It catches the
    mechanical defects that are invisible without a picture.

WHAT IT CHECKS (per slide)
    L1  off-canvas / clipped shapes      error when the shape carries text or
                                         sits wholly outside; warn for a
                                         decorative bleed
    L2  footer-band intrusion            warn  content below ~7.1in that is
                                         not footer chrome
    L3  text overflow risk               warn  past 1.5x the box, error past
                                         3x; estimated, see CALIBRATION
    L4  inherited font / colour / fill   warn  nothing explicit was set
    L5  theme-colour usage               warn  resolves via the Office theme
    L6  contrast below WCAG              warn  under 4.5:1 body / 3:1 large,
                                         error under 2:1
    L7  inconsistent inherited shadow    warn  once per deck
    L8  unresolvable brand font          error 'Fraunces Light', 'Fraunces
                                         9pt', bare 'Fraunces'
    L9  non-brand font                   warn

    L4, L5 and L8 matter because every FurtherAI template ships the **stock
    Office theme** — Arial, Office blue. Anything left inherited silently
    renders as that, not as the brand. See `assets/brand/BRAND.md`.

WHAT IT CHECKS (per deck, only with --type)
    T1  slide count off contract       warn  outside the type's documented range
    T2  wrong last slide               error the type ends on `next_steps` /
                                       `closing` and this deck does not

    A .pptx carries no record of the pattern each slide was built from, so T2
    reads the last slide's text: `next_steps` is recognised by its title,
    `closing` by carrying nothing but a logo and footer chrome. It stays quiet
    when it cannot tell. Structural conformance — required slots, slot order,
    sanctioned patterns — is checked on the spec instead, before the build:
    `build_deck.py spec.json --check --type NAME`.

CALIBRATION
    Thresholds were tuned against the hand-built decks in
    `assets/brand/pptx_templates/`, which are known-good by construction: a
    designer routinely undersizes a one-line label box and lets it spill
    harmlessly, bleeds artwork past the edge on purpose, and leaves shadow
    inheritance on everywhere. If a change to this file makes those decks
    light up, the change is wrong. They should come back near-clean apart
    from L8, which they genuinely trip.

    Two blind spots, both deliberate and both under-reporting rather than
    over-reporting:
      * contrast is skipped for text sitting on a picture — artwork has no
        knowable background colour;
      * overflow is a character-count estimate with no font metrics, so it
        catches gross breaks and misses marginal ones.

USAGE
    lint_deck.py <deck.pptx> [--type NAME] [--form NAME]
                 [--json] [--only L1,T2] [--quiet] [--max N]

EXIT CODES
    0   clean, or warnings only          — safe to ship
    1   at least one error-level finding — a build gating on this should stop
    2   bad usage, unreadable file, or python-pptx missing
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import deck_types as dtypes                                    # noqa: E402

try:
    from pptx import Presentation
    from pptx.util import Emu
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.enum.text import MSO_AUTO_SIZE
except ImportError:
    print("ERROR: needs python-pptx: python3 -m pip install python-pptx", file=sys.stderr)
    raise SystemExit(2)

# --------------------------------------------------------------- brand -----
# Family names that actually resolve from assets/brand/fonts/*/static/.
BRAND_FONTS = {
    "Wix Madefor Display", "Wix Madefor Display Medium",
    "Wix Madefor Display SemiBold", "Wix Madefor Display ExtraBold",
    "Fraunces 72pt", "Fraunces 72pt Light", "Fraunces 72pt SemiBold",
    "Courier New",
}
# Named in the templates but NOT resolvable — see BRAND.md "Fonts".
#   Fraunces Light / Fraunces 9pt  -> no TTF registers that family
#   Fraunces (bare)                -> variable font, default instance is Black 900
UNRESOLVABLE_FONTS = {"Fraunces Light", "Fraunces 9pt", "Fraunces"}

# Every deck here ships the stock Office theme, so theme slots resolve to these.
OFFICE_THEME = {
    "DARK_1": "000000", "LIGHT_1": "FFFFFF", "DARK_2": "44546A", "LIGHT_2": "E7E6E6",
    "ACCENT_1": "4472C4", "ACCENT_2": "ED7D31", "ACCENT_3": "A5A5A5",
    "ACCENT_4": "FFC000", "ACCENT_5": "5B9BD5", "ACCENT_6": "70AD47",
    "TEXT_1": "000000", "TEXT_2": "44546A", "BACKGROUND_1": "FFFFFF",
    "BACKGROUND_2": "E7E6E6", "HYPERLINK": "0563C1", "FOLLOWED_HYPERLINK": "954F72",
}

FOOTER_BAND_TOP = 7.10          # inches; below this is chrome on the 7.5in canvas
FOOTER_WORDS = re.compile(
    r"furtherai\.com|ai workspace for insurance|confidential|prepared for|^\d{1,2}$",
    re.I)

# T2 signatures. A .pptx does not record which pattern drew a slide, so the two
# closing patterns are recognised from what they put on the slide: next_steps
# titles itself, and a closing carries a logo plus footer chrome and nothing
# else. Anything else returns None and the check stays quiet.
NEXT_STEPS_RE = re.compile(r"next\s+steps?", re.I)
CLOSING_MAX_CHARS = 160

# Overflow model. Deliberately loose — a conservative flag beats none.
# Calibrated against the two hand-built brand decks, which are known-good: a
# designer routinely sets a one-line label box shorter than its own text and
# lets it spill harmlessly, so only real multi-line body copy is worth flagging.
CHAR_W = 0.48                   # mean advance as a fraction of point size
LINE_H = 1.22                   # line box as a multiple of point size
OVERFLOW_TOL = 1.50             # only complain past 50% over the box
OVERFLOW_ERR = 3.00             # ...and only call it an error past 3x
OVERFLOW_MIN_CHARS = 80         # below this it is a label, not a paragraph
DEFAULT_PT = 18.0               # PowerPoint's default when nothing is set


# ----------------------------------------------------------- utilities -----
def inches(v):
    return None if v is None else round(Emu(int(v)).inches, 3)


def _srgb(c):
    c = c / 255.0
    return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4


def luminance(hexrgb: str) -> float:
    r, g, b = (int(hexrgb[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * _srgb(r) + 0.7152 * _srgb(g) + 0.0722 * _srgb(b)


def contrast(fg: str, bg: str) -> float:
    a, b = luminance(fg), luminance(bg)
    hi, lo = max(a, b), min(a, b)
    return (hi + 0.05) / (lo + 0.05)


def color_hex(cf):
    """Resolve a ColorFormat to 'RRGGBB', or None when it is inherited."""
    if cf is None:
        return None
    try:
        t = str(cf.type)
    except Exception:
        return None
    if "RGB" in t:
        try:
            return str(cf.rgb)
        except Exception:
            return None
    if "SCHEME" in t or "THEME" in t:
        name = str(getattr(cf, "theme_color", "")).split(" ")[0]
        return OFFICE_THEME.get(name)
    return None


def is_theme_color(cf) -> bool:
    try:
        return cf is not None and "SCHEME" in str(cf.type)
    except Exception:
        return False


def shape_fill_hex(shape):
    """Solid fill hex, 'none' when explicitly no fill, None when inherited."""
    try:
        f = shape.fill
        t = str(f.type)
    except Exception:
        return None
    if "BACKGROUND" in t:
        return "none"
    if "SOLID" in t:
        return color_hex(f.fore_color)
    if "GRADIENT" in t:
        try:
            return color_hex(f.gradient_stops[0].color)
        except Exception:
            return None
    return None                      # patterned, picture, or inherited


def slide_bg_hex(slide, prs):
    for src in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        try:
            f = src.background.fill
            t = str(f.type)
        except Exception:
            continue
        if "SOLID" in t:
            h = color_hex(f.fore_color)
            if h:
                return h
        if "GRADIENT" in t:
            try:
                h = color_hex(f.gradient_stops[0].color)
                if h:
                    return h
            except Exception:
                pass
    return None


def has_fullbleed_image(slide, cw, ch):
    """A full-bleed picture makes every contrast check downstream a guess."""
    for sh in slide.shapes:
        if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        try:
            w, h = inches(sh.width) or 0, inches(sh.height) or 0
        except Exception:
            continue
        if w >= cw * 0.8 and h >= ch * 0.8:
            return True
    return False


def fill_layers(slide):
    """Top-level solid-filled shapes in z-order (document order = back to front).

    These decks put an unfilled text box on top of a separately drawn coloured
    rectangle, so the visual background of a run is the nearest filled shape
    *underneath it*, never the slide background. Resolving against the slide
    background instead reports white-on-white for every card header.
    """
    layers = []
    for i, sh in enumerate(slide.shapes):
        if sh.shape_type in (MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.PICTURE):
            continue
        hexv = shape_fill_hex(sh)
        if not hexv or hexv == "none":
            continue
        L, T, W, H = inches(sh.left), inches(sh.top), inches(sh.width), inches(sh.height)
        if None in (L, T, W, H):
            continue
        layers.append((i, L, T, L + W, T + H, hexv))
    return layers


def picture_boxes(slide):
    """Top-level pictures in z-order. Text sitting on artwork has no knowable
    background colour, so contrast must not be guessed there."""
    out = []
    for i, sh in enumerate(slide.shapes):
        if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        L, T, W, H = inches(sh.left), inches(sh.top), inches(sh.width), inches(sh.height)
        if None in (L, T, W, H):
            continue
        out.append((i, L, T, L + W, T + H))
    return out


def over_picture(shape, z_index, pics) -> bool:
    L, T, W, H = inches(shape.left), inches(shape.top), inches(shape.width), inches(shape.height)
    if None in (L, T, W, H):
        return False
    cx, cy = L + W / 2.0, T + H / 2.0
    return any(i < z_index and x0 - 0.01 <= cx <= x1 + 0.01 and y0 - 0.01 <= cy <= y1 + 0.01
               for i, x0, y0, x1, y1 in pics)


def bg_under(shape, z_index, layers, slide_bg):
    """Nearest filled shape strictly below `shape` in z-order that covers its
    centre; falls back to the slide background."""
    L, T, W, H = inches(shape.left), inches(shape.top), inches(shape.width), inches(shape.height)
    if None in (L, T, W, H):
        return slide_bg
    cx, cy = L + W / 2.0, T + H / 2.0
    for i, x0, y0, x1, y1, hexv in reversed(layers):
        if i >= z_index:
            continue
        if x0 - 0.01 <= cx <= x1 + 0.01 and y0 - 0.01 <= cy <= y1 + 0.01:
            return hexv
    return slide_bg


def run_pt(run, para, default=None):
    for src in (run.font, para.font):
        try:
            if src.size is not None:
                return round(src.size.pt, 1)
        except Exception:
            pass
    return default


def para_text(para):
    return "".join(r.text for r in para.runs)


# -------------------------------------------------------------- checks -----
def estimate_height_in(shape) -> float | None:
    """Estimated rendered text height, inches. None when not estimable."""
    tf = shape.text_frame
    try:
        ml = Emu(int(tf.margin_left)).inches
        mr = Emu(int(tf.margin_right)).inches
        mt = Emu(int(tf.margin_top)).inches
        mb = Emu(int(tf.margin_bottom)).inches
    except Exception:
        ml = mr = 0.1
        mt = mb = 0.05
    avail_w = (inches(shape.width) or 0) - ml - mr
    if avail_w <= 0.05:
        return None
    total = 0.0
    for para in tf.paragraphs:
        txt = para_text(para)
        if not txt.strip():
            total += DEFAULT_PT * LINE_H / 72.0
            continue
        pt = None
        for r in para.runs:
            pt = run_pt(r, para)
            if pt:
                break
        if pt is None:
            return None                       # inherited size — don't guess
        cpl = max(1, int(avail_w * 72.0 / (CHAR_W * pt)))
        lines = 1 if tf.word_wrap is False else max(1, math.ceil(len(txt) / cpl))
        ls = 1.0
        try:
            if isinstance(para.line_spacing, float):
                ls = para.line_spacing
            elif para.line_spacing is not None:
                ls = para.line_spacing.pt / pt
        except Exception:
            pass
        total += lines * pt * LINE_H * ls / 72.0
        for attr in ("space_before", "space_after"):
            try:
                v = getattr(para, attr)
                if v is not None:
                    total += v.pt / 72.0
            except Exception:
                pass
    return total + mt + mb


def walk(shapes, z=None, depth=0):
    """Yield (shape, top_level_z_index or None). Group children have their own
    coordinate space, so only top-level shapes carry a z index and get the
    geometry checks."""
    for i, sh in enumerate(shapes):
        zi = i if depth == 0 else z
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield sh, (zi if depth == 0 else None)
            yield from walk(sh.shapes, zi, depth + 1)
        else:
            yield sh, (zi if depth == 0 else None)


def content_text(slide):
    """Text on a slide with the footer chrome dropped, longest run first."""
    out = []
    for shape, _z in walk(slide.shapes):
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        if not text:
            continue
        if len(text) < 60 and FOOTER_WORDS.search(text):
            continue
        out.append(text)
    return sorted(out, key=len, reverse=True)


def reads_as(pattern, texts):
    """True / False / None -- whether a slide reads as `pattern`.

    None means no signature is known for that pattern, so nothing is claimed.
    """
    if pattern == "next_steps":
        return any(NEXT_STEPS_RE.search(t) for t in texts)
    if pattern == "closing":
        return len(texts) <= 1 and sum(len(t) for t in texts) <= CLOSING_MAX_CHARS
    return None


def lint(path: Path, only=None, type_name=None, form=None):
    prs = Presentation(str(path))
    cw = round(Emu(prs.slide_width).inches, 3)
    ch = round(Emu(prs.slide_height).inches, 3)
    out = []

    def add(code, sev, s_i, loc, msg):
        if only and code not in only:
            return
        out.append({"code": code, "sev": sev, "slide": s_i, "loc": loc, "msg": msg})

    shadow_inherit = 0
    shadow_total = 0

    for s_i, slide in enumerate(prs.slides, start=1):
        bg = slide_bg_hex(slide, prs)
        murky = has_fullbleed_image(slide, cw, ch)
        layers = fill_layers(slide)
        pics = picture_boxes(slide)
        seen = set()                      # per-slide dedupe for style findings

        def once(key):
            if key in seen:
                return False
            seen.add(key)
            return True

        for sh_i, (shape, z) in enumerate(walk(slide.shapes)):
            name = getattr(shape, "name", "?")
            loc = f"s:{s_i}/sh:{sh_i}"
            top = z is not None
            L, T = inches(shape.left), inches(shape.top)
            W, H = inches(shape.width), inches(shape.height)
            has_text = shape.has_text_frame and shape.text_frame.text.strip()

            # -- L1 off-canvas -------------------------------------------
            # A designer bleeds decorative art past the edge on purpose, so a
            # shape that still overlaps the canvas is only a defect when it
            # carries text (i.e. content is actually being clipped).
            if top and None not in (L, T, W, H):
                R, B = L + W, T + H
                over = []
                if L < -0.02:
                    over.append(f"left {L}in")
                if T < -0.02:
                    over.append(f"top {T}in")
                if R > cw + 0.02:
                    over.append(f"right edge {round(R,2)}in > {cw}in")
                if B > ch + 0.02:
                    over.append(f"bottom edge {round(B,2)}in > {ch}in")
                fully_out = R <= 0.02 or B <= 0.02 or L >= cw - 0.02 or T >= ch - 0.02
                if over and (has_text or fully_out):
                    add("L1", "error", s_i, loc,
                        f"{name!r} {'entirely off-canvas' if fully_out else 'clipped'}: "
                        f"{', '.join(over)}")
                elif over:
                    add("L1", "warn", s_i, loc,
                        f"{name!r} bleeds past the canvas ({', '.join(over)}) — "
                        f"decorative, no text")
                # -- L2 footer band --------------------------------------
                elif T >= FOOTER_BAND_TOP and has_text:
                    txt = shape.text_frame.text.strip()
                    if not FOOTER_WORDS.search(txt):
                        add("L2", "warn", s_i, loc,
                            f"{name!r} sits in the footer band (top {T}in): {txt[:40]!r}")

            # -- L7 inherited shadow (counted, reported once per deck) ----
            try:
                shadow_total += 1
                if shape.shadow.inherit:
                    shadow_inherit += 1
            except Exception:
                shadow_total -= 1

            # -- L4 inherited fill ---------------------------------------
            # Connectors, straight lines and tables carry no meaningful fill;
            # only closed shapes that paint a background are worth checking.
            fillable = shape.shape_type not in (
                MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.LINE,
                MSO_SHAPE_TYPE.TABLE,
            ) and not getattr(shape, "is_placeholder", False)
            if fillable:
                if shape_fill_hex(shape) is None and once(("fill", s_i)):
                    add("L4", "warn", s_i, loc,
                        f"shape fill inherited rather than explicit "
                        f"(first: {name!r}) — theme default applies")

            if not shape.has_text_frame:
                continue

            # -- L3 overflow ---------------------------------------------
            auto = None
            try:
                auto = shape.text_frame.auto_size
            except Exception:
                pass
            grows = auto == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
            body = shape.text_frame.text.strip()
            if top and H and len(body) >= OVERFLOW_MIN_CHARS and not grows:
                need = estimate_height_in(shape)
                if need is not None and need > H * OVERFLOW_TOL:
                    ratio = need / H
                    sev = "error" if ratio > OVERFLOW_ERR else "warn"
                    add("L3", sev, s_i, loc,
                        f"{name!r} text needs ~{need:.2f}in in a {H}in box "
                        f"({ratio:.1f}x): {body[:44]!r}")

            # -- per-run checks ------------------------------------------
            own = shape_fill_hex(shape)
            on_art = top and over_picture(shape, z, pics)
            if own and own != "none":
                eff_bg = own
            elif top and not on_art:
                eff_bg = bg_under(shape, z, layers, bg)
            elif on_art:
                eff_bg = None          # artwork behind it — unknowable
            else:
                eff_bg = bg
            for p_i, para in enumerate(shape.text_frame.paragraphs):
                for r in para.runs:
                    if not r.text.strip():
                        continue
                    rloc = f"{loc}/p:{p_i}"
                    fam = r.font.name
                    if fam is None:
                        if once(("nofont", s_i)):
                            add("L4", "warn", s_i, rloc,
                                f"run has no explicit font, e.g. {r.text.strip()[:34]!r}")
                    elif fam in UNRESOLVABLE_FONTS:
                        if once(("badfont", s_i, fam)):
                            add("L8", "error", s_i, rloc,
                                f"{fam!r} does not resolve (see BRAND.md) — use "
                                f"{'Fraunces 72pt Light' if 'Light' in fam else 'Fraunces 72pt'}")
                    elif fam not in BRAND_FONTS:
                        if once(("offbrand", s_i, fam)):
                            add("L9", "warn", s_i, rloc, f"non-brand font {fam!r}")

                    cf = r.font.color
                    if is_theme_color(cf) and once(("theme", s_i)):
                        add("L5", "warn", s_i, rloc,
                            f"theme color {getattr(cf,'theme_color','?')} "
                            f"(resolves via the stock Office theme), e.g. {r.text.strip()[:24]!r}")
                    fg = color_hex(cf)
                    if fg is None and once(("nocolor", s_i)):
                        add("L4", "warn", s_i, rloc,
                            f"run has no explicit color, e.g. {r.text.strip()[:34]!r}")

                    # -- L6 contrast -------------------------------------
                    if fg and eff_bg and eff_bg != "none" and not murky and not on_art:
                        pt = run_pt(r, para, DEFAULT_PT)
                        large = pt >= 18 or (pt >= 14 and bool(r.font.bold))
                        need = 3.0 if large else 4.5
                        ratio = contrast(fg, eff_bg)
                        if ratio < need:
                            sev = "error" if ratio < 2.0 else "warn"
                            add("L6", sev, s_i, rloc,
                                f"#{fg} on #{eff_bg} = {ratio:.2f}:1 "
                                f"(needs {need}:1 at {pt}pt): {r.text.strip()[:30]!r}")

    if shadow_total and shadow_inherit:
        pct = 100.0 * shadow_inherit / shadow_total
        # Inheriting is the python-pptx default and the stock theme draws no
        # shadow, so this only matters when a build sets it on some shapes and
        # not others. A deck that inherits everywhere renders fine.
        if pct < 95.0:
            add("L7", "warn", 0, "deck",
                f"{shadow_inherit}/{shadow_total} shapes ({pct:.0f}%) inherit a drop "
                f"shadow while others set one explicitly — inconsistent")

    n_slides = len(prs.slides)
    if type_name:
        lo, hi = dtypes.length_range(type_name, form)
        if not lo <= n_slides <= hi:
            add("T1", "warn", 0, "deck",
                f"{n_slides} slides; {type_name} is documented at "
                f"{lo if lo == hi else f'{lo}-{hi}'} — a repeating slot can "
                f"stretch that, a thin deck cannot")
        want = dtypes.last_pattern(type_name, form)
        if want and n_slides:
            last = prs.slides[n_slides - 1]
            texts = content_text(last)
            verdict = reads_as(want, texts)
            if verdict is False:
                sample = (texts[0][:44] if texts else "(no text)")
                add("T2", "error", n_slides, f"s:{n_slides}",
                    f"{type_name} ends on `{want}` — this last slide does not "
                    f"read as one: {sample!r}")
    return out, cw, ch, n_slides


# -------------------------------------------------------------- report -----
LABEL = {
    "L1": "off-canvas", "L2": "footer band", "L3": "overflow risk",
    "L4": "inherited style", "L5": "theme color", "L6": "contrast",
    "L7": "inherited shadow", "L8": "unresolvable font", "L9": "non-brand font",
    "T1": "type slide count", "T2": "type last slide",
}


def render(findings, path, cw, ch, n_slides, quiet, max_per):
    errs = [f for f in findings if f["sev"] == "error"]
    warns = [f for f in findings if f["sev"] == "warn"]
    lines = [f"# Deck lint — {Path(path).name}",
             f"{n_slides} slides on {cw} x {ch}in · "
             f"**{len(errs)} error{'' if len(errs)==1 else 's'}, "
             f"{len(warns)} warning{'' if len(warns)==1 else 's'}**", ""]
    by_code = {}
    for f in findings:
        by_code.setdefault(f["code"], []).append(f)
    lines.append("| Check | Errors | Warnings |")
    lines.append("|---|---|---|")
    for code in sorted(by_code, key=lambda c: (-sum(1 for x in by_code[c] if x["sev"] == "error"), c)):
        g = by_code[code]
        e = sum(1 for x in g if x["sev"] == "error")
        w = len(g) - e
        lines.append(f"| {code} {LABEL[code]} | {e or ''} | {w or ''} |")
    if quiet:
        return "\n".join(lines)
    lines.append("")
    cur = None
    shown = {}
    for f in sorted(findings, key=lambda f: (f["slide"], 0 if f["sev"] == "error" else 1, f["code"])):
        if f["slide"] != cur:
            cur = f["slide"]
            lines.append(f"\n**Slide {cur}**")
            shown = {}
        k = f["code"]
        shown[k] = shown.get(k, 0) + 1
        if shown[k] > max_per:
            if shown[k] == max_per + 1:
                lines.append(f"  … more {k} on this slide")
            continue
        mark = "ERR " if f["sev"] == "error" else "warn"
        lines.append(f"  {mark} [{f['code']}] {f['loc']}  {f['msg']}")
    return "\n".join(lines)


def main() -> int:
    ap = argparse.ArgumentParser(
        prog="lint_deck.py",
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="exit 0 = clean or warnings only · 1 = error-level finding "
               "(gate on this) · 2 = usage or read error",
    )
    ap.add_argument("file", help="the .pptx to lint")
    ap.add_argument("--type", metavar="NAME",
                    help="also check the deck against a type contract "
                         "(see build_deck.py --list-types)")
    ap.add_argument("--form", metavar="NAME",
                    help="which form of a multi-form type, e.g. "
                         "--type implementation_checkin --form tracker_only")
    ap.add_argument("--json", action="store_true",
                    help="emit findings as JSON for machine consumption")
    ap.add_argument("--only", metavar="CODES",
                    help="run only these checks, comma-separated, e.g. L1,T2")
    ap.add_argument("--quiet", action="store_true",
                    help="summary table only, no per-slide detail")
    ap.add_argument("--max", type=int, default=4, metavar="N",
                    help="cap findings shown per check per slide (default 4)")
    a = ap.parse_args()

    path = Path(a.file).expanduser()
    if not path.exists():
        print(f"ERROR: no such file: {path}", file=sys.stderr)
        return 2
    if path.suffix.lower() != ".pptx":
        print(f"ERROR: expected a .pptx, got {path.suffix or '(no extension)'}",
              file=sys.stderr)
        return 2
    if a.form and not a.type:
        print("ERROR: --form needs --type NAME", file=sys.stderr)
        return 2
    if a.type:
        try:
            dtypes.slot_plan(a.type, a.form)
        except (dtypes.UnknownType, dtypes.UnknownForm) as exc:
            print(f"ERROR: {exc}", file=sys.stderr)
            return 2
    known = set(LABEL)
    only = None
    if a.only:
        only = {c.strip().upper() for c in a.only.split(",") if c.strip()}
        unknown = only - known
        if unknown:
            print(f"ERROR: unknown check(s): {', '.join(sorted(unknown))} — "
                  f"known: {', '.join(sorted(known))}", file=sys.stderr)
            return 2
    try:
        findings, cw, ch, n = lint(path, only, a.type, a.form)
    except Exception as exc:                       # a malformed deck, not a crash
        print(f"ERROR: could not read {path.name}: {exc}", file=sys.stderr)
        return 2

    if a.json:
        payload = {"file": str(path), "canvas": [cw, ch], "slides": n}
        if a.type:
            payload["type"] = a.type
            payload["form"] = a.form or dtypes.default_form(a.type)
        payload["findings"] = findings
        print(json.dumps(payload, indent=2))
    else:
        print(render(findings, path, cw, ch, n, a.quiet, a.max))
    return 1 if any(f["sev"] == "error" for f in findings) else 0


if __name__ == "__main__":
    raise SystemExit(main())
