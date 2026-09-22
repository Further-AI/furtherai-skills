#!/usr/bin/env python3
"""inspect_docx.py — dump the structure of a .docx so a new document can be
mapped onto it.

Reports page setup, document defaults, embedded fonts, header/footer parts, and
then every body block in document order: paragraphs with their style, spacing,
borders and per-run typography; tables as grids with per-cell width, fill,
margins, borders and text; inline images with their placed size.

    python3 inspect_docx.py FILE.docx            # full dump
    python3 inspect_docx.py FILE.docx --text-only # just the text, in order
    python3 inspect_docx.py FILE.docx --json      # machine-readable

Requires python-docx (``python3 -m pip install python-docx``).
"""
from __future__ import annotations

import argparse
import json
import sys
import zipfile
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:  # pragma: no cover - dependency guard
    sys.stderr.write(
        "inspect_docx requires python-docx.\n"
        "  python3 -m pip install python-docx\n"
    )
    raise SystemExit(1)

EMU_PER_INCH = 914400
TWIPS_PER_INCH = 1440


def _in(value, per=EMU_PER_INCH):
    return None if value is None else round(value / per, 3)


def _attr(el, name):
    return None if el is None else el.get(qn("w:" + name))


def _border(el):
    if el is None:
        return None
    out = {"val": _attr(el, "val")}
    for name in ("sz", "space", "color"):
        value = _attr(el, name)
        if value is not None:
            out[name] = value
    return out


def _borders(parent, tag, edges):
    if parent is None:
        return None
    node = parent.find(qn("w:" + tag))
    if node is None:
        return None
    out = {}
    for edge in edges:
        border = _border(node.find(qn("w:" + edge)))
        if border:
            out[edge] = border
    return out or None


def _shd_fill(tc_pr):
    if tc_pr is None:
        return None
    shd = tc_pr.find(qn("w:shd"))
    fill = _attr(shd, "fill")
    return None if fill in (None, "auto") else fill.upper()


def _tc_margins(tc_pr):
    if tc_pr is None:
        return None
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        return None
    out = {}
    for edge in ("top", "left", "bottom", "right"):
        node = mar.find(qn("w:" + edge))
        if node is not None:
            out[edge] = int(float(node.get(qn("w:w"))))
    return out or None


def _run_info(run):
    rpr = run._element.find(qn("w:rPr"))
    rfonts = None if rpr is None else rpr.find(qn("w:rFonts"))
    sz = None if rpr is None else rpr.find(qn("w:sz"))
    color = None if rpr is None else rpr.find(qn("w:color"))
    bold = None if rpr is None else rpr.find(qn("w:b"))
    italic = None if rpr is None else rpr.find(qn("w:i"))

    def flag(node):
        if node is None:
            return False
        return _attr(node, "val") not in ("0", "false")

    info = {
        "text": run.text,
        "font": _attr(rfonts, "ascii"),
        "font_hAnsi": _attr(rfonts, "hAnsi"),
        "font_cs": _attr(rfonts, "cs"),
        "size_pt": (float(_attr(sz, "val")) / 2) if sz is not None else None,
        "color": (_attr(color, "val") or "").upper() or None,
        "bold": flag(bold),
        "italic": flag(italic),
    }
    images = []
    for drawing in run._element.findall(".//" + qn("w:drawing")):
        extent = drawing.find(".//" + qn("wp:extent"))
        blip = drawing.find(".//" + qn("a:blip"))
        images.append({
            "width_in": _in(int(extent.get("cx"))) if extent is not None else None,
            "height_in": _in(int(extent.get("cy"))) if extent is not None else None,
            "rel_id": blip.get(qn("r:embed")) if blip is not None else None,
        })
    if images:
        info["images"] = images
    return info


def _paragraph_info(paragraph, part=None):
    pf = paragraph.paragraph_format
    ppr = paragraph._p.find(qn("w:pPr"))
    info = {
        "kind": "paragraph",
        "style": paragraph.style.name if paragraph.style is not None else None,
        "text": paragraph.text,
        "align": str(pf.alignment) if pf.alignment is not None else None,
        "space_before_pt": _in(pf.space_before, 12700) if pf.space_before is not None else None,
        "space_after_pt": _in(pf.space_after, 12700) if pf.space_after is not None else None,
        "line_spacing": pf.line_spacing,
        "left_indent_in": _in(pf.left_indent) if pf.left_indent is not None else None,
        "borders": _borders(ppr, "pBdr", ("top", "left", "bottom", "right", "between")),
        "runs": [_run_info(r) for r in paragraph.runs],
    }
    images = []
    for run in info["runs"]:
        images.extend(run.pop("images", []))
    if part is not None:
        for image in images:
            rel_id = image.get("rel_id")
            if rel_id and rel_id in part.rels:
                image["target"] = part.rels[rel_id].target_ref
    if images:
        info["images"] = images
    return info


def _table_info(table, part=None):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    grid = [int(float(c.get(qn("w:w"))))
            for c in tbl.findall(qn("w:tblGrid") + "/" + qn("w:gridCol"))]
    tblw = None if tbl_pr is None else tbl_pr.find(qn("w:tblW"))
    layout = None if tbl_pr is None else tbl_pr.find(qn("w:tblLayout"))
    look = None if tbl_pr is None else tbl_pr.find(qn("w:tblLook"))
    style = None if tbl_pr is None else tbl_pr.find(qn("w:tblStyle"))

    info = {
        "kind": "table",
        "rows": len(table.rows),
        "cols": len(table.columns),
        "style": _attr(style, "val"),
        "width_tw": int(float(_attr(tblw, "w"))) if tblw is not None else None,
        "layout": _attr(layout, "type"),
        "look": _attr(look, "val"),
        "grid_tw": grid,
        "grid_in": [round(g / TWIPS_PER_INCH, 3) for g in grid],
        "borders": _borders(tbl_pr, "tblBorders",
                            ("top", "left", "bottom", "right", "insideH", "insideV")),
        "cells": [],
    }
    for r_idx, row in enumerate(table.rows):
        tr_pr = row._tr.find(qn("w:trPr"))
        header_el = None if tr_pr is None else tr_pr.find(qn("w:tblHeader"))
        header = header_el is not None and _attr(header_el, "val") not in ("0", "false")
        for c_idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.find(qn("w:tcPr"))
            tcw = None if tc_pr is None else tc_pr.find(qn("w:tcW"))
            valign = None if tc_pr is None else tc_pr.find(qn("w:vAlign"))
            info["cells"].append({
                "row": r_idx,
                "col": c_idx,
                "header_row": header,
                "width_tw": int(float(_attr(tcw, "w"))) if tcw is not None else None,
                "fill": _shd_fill(tc_pr),
                "margins_tw": _tc_margins(tc_pr),
                "borders": _borders(tc_pr, "tcBorders", ("top", "left", "bottom", "right")),
                "valign": _attr(valign, "val"),
                "paragraphs": [_paragraph_info(p, part) for p in cell.paragraphs],
            })
    return info


def _doc_defaults(doc):
    styles = doc.styles.element
    docdefaults = styles.find(qn("w:docDefaults"))
    if docdefaults is None:
        return {}
    rpr = docdefaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
    if rpr is None:
        return {}
    rfonts = rpr.find(qn("w:rFonts"))
    sz = rpr.find(qn("w:sz"))
    color = rpr.find(qn("w:color"))
    return {
        "font": _attr(rfonts, "ascii"),
        "size_pt": (float(_attr(sz, "val")) / 2) if sz is not None else None,
        "color": (_attr(color, "val") or "").upper() or None,
    }


def _embedded_fonts(path):
    import re

    out = {"parts": [], "fontTable": []}
    with zipfile.ZipFile(path) as zf:
        out["parts"] = sorted(n for n in zf.namelist() if n.startswith("word/fonts/"))
        try:
            xml = zf.read("word/fontTable.xml").decode("utf-8")
        except KeyError:
            return out
    for match in re.finditer(r'<w:font w:name="([^"]+)">(.*?)</w:font>', xml, re.S):
        name, block = match.groups()
        styles = re.findall(r"<w:embed(\w+)\s", block)
        if styles:
            out["fontTable"].append({"family": name, "embeds": styles})
    return out


def _section_info(section):
    return {
        "page_width_in": _in(section.page_width),
        "page_height_in": _in(section.page_height),
        "orientation": "portrait" if str(section.orientation).startswith("PORTRAIT") else "landscape",
        "margins_in": {
            "left": _in(section.left_margin),
            "right": _in(section.right_margin),
            "top": _in(section.top_margin),
            "bottom": _in(section.bottom_margin),
        },
        "header_distance_in": _in(section.header_distance),
        "footer_distance_in": _in(section.footer_distance),
    }


def _part_paragraphs(container, part):
    return [_paragraph_info(p, part) for p in container.paragraphs]


def inspect(path):
    """Parse a .docx into a plain-dict description."""
    path = Path(path)
    doc = Document(str(path))
    body = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph

            body.append(_paragraph_info(Paragraph(child, doc), doc.part))
        elif child.tag == qn("w:tbl"):
            from docx.table import Table

            body.append(_table_info(Table(child, doc), doc.part))
    sections = []
    for section in doc.sections:
        info = _section_info(section)
        info["header"] = {
            "linked": section.header.is_linked_to_previous,
            "paragraphs": _part_paragraphs(section.header, doc.part),
        }
        info["footer"] = {
            "linked": section.footer.is_linked_to_previous,
            "paragraphs": _part_paragraphs(section.footer, doc.part),
        }
        sections.append(info)
    return {
        "file": str(path),
        "sections": sections,
        "defaults": _doc_defaults(doc),
        "fonts": _embedded_fonts(path),
        "styles": sorted(s.name for s in doc.styles if s.name),
        "body": body,
    }


# --------------------------------------------------------------------------
# Text rendering
# --------------------------------------------------------------------------


def _fmt_border(border):
    if not border:
        return "-"
    parts = [border.get("val", "?")]
    if border.get("sz"):
        parts.append("sz" + str(border["sz"]))
    if border.get("color") and border["color"] not in ("auto",):
        parts.append("#" + str(border["color"]).upper())
    return " ".join(parts)


def _fmt_borders(borders):
    if not borders:
        return "(none)"
    return "  ".join("{}={}".format(k, _fmt_border(v)) for k, v in borders.items())


def _fmt_run(run):
    bits = []
    if run.get("font"):
        bits.append("'{}'".format(run["font"]))
    if run.get("size_pt") is not None:
        bits.append("{:g}pt".format(run["size_pt"]))
    if run.get("color"):
        bits.append("#" + run["color"])
    if run.get("bold"):
        bits.append("bold")
    if run.get("italic"):
        bits.append("italic")
    return " ".join(bits) or "(inherited)"


def _fmt_paragraph(info, indent, out):
    pad = " " * indent
    bits = ["style={}".format(info["style"])]
    if info["space_before_pt"]:
        bits.append("before={:g}pt".format(info["space_before_pt"]))
    if info["space_after_pt"] is not None:
        bits.append("after={:g}pt".format(info["space_after_pt"]))
    if info["line_spacing"]:
        bits.append("line={}".format(info["line_spacing"]))
    if info["align"]:
        bits.append("align={}".format(info["align"]))
    if info["left_indent_in"]:
        bits.append("indent={}in".format(info["left_indent_in"]))
    out.append(pad + "P  " + "  ".join(bits))
    if info["borders"]:
        out.append(pad + "   borders: " + _fmt_borders(info["borders"]))
    for image in info.get("images", []):
        out.append(pad + "   image {} {}x{}in".format(
            image.get("target", image.get("rel_id")),
            image["width_in"], image["height_in"]))
    for run in info["runs"]:
        if not run["text"]:
            continue
        out.append(pad + '   "{}"'.format(run["text"]) + "  [" + _fmt_run(run) + "]")


def render(data, text_only=False):
    out = []
    if text_only:
        for block in data["body"]:
            if block["kind"] == "paragraph":
                if block["text"].strip():
                    out.append(block["text"])
            else:
                grid = {}
                for cell in block["cells"]:
                    grid.setdefault(cell["row"], []).append(
                        " ".join(p["text"] for p in cell["paragraphs"]).strip())
                out.append("TABLE {}x{}".format(block["rows"], block["cols"]))
                for r_idx in sorted(grid):
                    out.append("  | " + " | ".join(grid[r_idx]) + " |")
        return "\n".join(out)

    out.append("FILE      {}".format(data["file"]))
    for idx, section in enumerate(data["sections"]):
        margins = section["margins_in"]
        out.append(
            "PAGE  [{}] {} x {}in {} | margins L{} R{} T{} B{} | header {}in footer {}in".format(
                idx, section["page_width_in"], section["page_height_in"],
                section["orientation"], margins["left"], margins["right"],
                margins["top"], margins["bottom"],
                section["header_distance_in"], section["footer_distance_in"]))
        for part in ("header", "footer"):
            texts = [p["text"] for p in section[part]["paragraphs"] if p["text"].strip()]
            out.append("  {:<7} {}".format(
                part.upper(), " / ".join(texts) if texts else "(empty)"))
    defaults = data["defaults"]
    out.append("DEFAULTS  font={} size={}pt color={}".format(
        defaults.get("font"), defaults.get("size_pt"), defaults.get("color")))
    fonts = data["fonts"]
    if fonts["fontTable"]:
        for entry in fonts["fontTable"]:
            out.append("EMBEDDED  {} -> {}".format(entry["family"], ", ".join(entry["embeds"])))
        out.append("          {} font parts in word/fonts/".format(len(fonts["parts"])))
    else:
        out.append("EMBEDDED  (no fonts embedded)")
    out.append("")
    out.append("BODY ({} blocks)".format(len(data["body"])))
    for idx, block in enumerate(data["body"]):
        out.append("[{}]".format(idx))
        if block["kind"] == "paragraph":
            _fmt_paragraph(block, 4, out)
        else:
            out.append("    TABLE {}x{}  width={}tw ({:.2f}in)  layout={}  look={}  style={}".format(
                block["rows"], block["cols"], block["width_tw"],
                (block["width_tw"] or 0) / TWIPS_PER_INCH,
                block["layout"], block["look"], block["style"]))
            out.append("      grid: {} tw = {} in".format(
                " / ".join(str(g) for g in block["grid_tw"]),
                " / ".join(str(g) for g in block["grid_in"])))
            out.append("      borders: " + _fmt_borders(block["borders"]))
            for cell in block["cells"]:
                head = "      r{}c{} w={} fill={} margins={}".format(
                    cell["row"], cell["col"],
                    "{}tw".format(cell["width_tw"]) if cell["width_tw"] else "-",
                    cell["fill"] or "-", cell["margins_tw"] or "-")
                if cell["header_row"]:
                    head += " [header row]"
                if cell["borders"]:
                    head += " borders=" + _fmt_borders(cell["borders"])
                out.append(head)
                for paragraph in cell["paragraphs"]:
                    _fmt_paragraph(paragraph, 10, out)
    return "\n".join(out)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("docx", help="path to a .docx file")
    parser.add_argument("--text-only", action="store_true",
                        help="print only the document text, in reading order")
    parser.add_argument("--json", action="store_true", help="emit the full structure as JSON")
    args = parser.parse_args(argv)

    path = Path(args.docx)
    if not path.exists():
        sys.stderr.write("no such file: {}\n".format(path))
        return 1
    data = inspect(path)
    if args.json:
        print(json.dumps(data, indent=2))
    else:
        print(render(data, text_only=args.text_only))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
