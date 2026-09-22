#!/usr/bin/env python3
"""Dump an existing .pptx so new content can be mapped onto its layout.

Per slide: every shape with its index, name, type, inches geometry, fill, line,
and the text with run-level font / size / colour.  Tables print as grids,
pictures report their pixel dimensions.

    python3 inspect_pptx.py DECK.pptx                # everything
    python3 inspect_pptx.py DECK.pptx --slide 3      # one slide
    python3 inspect_pptx.py DECK.pptx --text-only    # just the copy
    python3 inspect_pptx.py DECK.pptx --json         # machine readable

Groups are recursed into and printed with a dotted index (e.g. 12.3).
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.oxml.ns import qn
except ImportError:  # pragma: no cover
    raise SystemExit("python-pptx is required.\n  python3 -m pip install python-pptx")

EMU = 914400.0


def _in(v):
    return None if v is None else round(v / EMU, 3)


def _fmt_in(v):
    return "  -  " if v is None else "%5.2f" % v


def _color_of(color_format):
    """'#RRGGBB' | 'scheme:ACCENT_1' | None -- never guesses."""
    try:
        ctype = color_format.type
    except (AttributeError, ValueError):
        return None
    if ctype is None:
        return None
    name = str(ctype).split(" ")[0].split(".")[-1]
    if name in ("RGB", "MSO_COLOR_TYPE.RGB", "1"):
        try:
            return "#%s" % str(color_format.rgb)
        except (AttributeError, ValueError):
            return None
    try:
        return "scheme:%s" % str(color_format.theme_color).split(" ")[0]
    except (AttributeError, ValueError):
        return name


def _grad_stops(shape):
    try:
        spPr = shape._element.spPr
    except AttributeError:
        return None
    grad = spPr.find(qn("a:gradFill")) if spPr is not None else None
    if grad is None:
        return None
    stops = []
    gs_lst = grad.find(qn("a:gsLst"))
    for gs in (gs_lst if gs_lst is not None else []):
        pos = int(gs.get("pos", "0")) / 100000.0
        clr = gs.find(qn("a:srgbClr"))
        stops.append((round(pos, 3), "#%s" % clr.get("val") if clr is not None else "?"))
    return stops or None


def _fill_of(shape):
    stops = _grad_stops(shape)
    if stops:
        return "grad(%s)" % " ".join("%.2f:%s" % (p, c) for p, c in stops)
    try:
        fill = shape.fill
        ftype = fill.type
    except (AttributeError, ValueError, NotImplementedError):
        return None
    if ftype is None:
        return "inherit"
    name = str(ftype).split(" ")[0].split(".")[-1]
    if name in ("SOLID", "1"):
        return _color_of(fill.fore_color) or "solid"
    if name in ("BACKGROUND", "5"):
        return "none"
    return name.lower()


def _line_of(shape):
    try:
        line = shape.line
    except (AttributeError, ValueError):
        return None
    try:
        if line.fill.type is not None and str(line.fill.type).startswith("BACKGROUND"):
            return "none"
    except (AttributeError, ValueError, NotImplementedError):
        pass
    color = _color_of(line.color)
    width = _in(line.width) if line.width else None
    if color is None and width is None:
        return None
    return "%s/%s" % (color or "?", ("%.3f" % width) if width else "?")


def _runs_of(text_frame):
    out = []
    for p in text_frame.paragraphs:
        for r in p.runs:
            if not r.text:
                continue
            f = r.font
            typeface = f.name
            if typeface is None:
                rPr = r._r.find(qn("a:rPr"))
                latin = rPr.find(qn("a:latin")) if rPr is not None else None
                typeface = latin.get("typeface") if latin is not None else None
            out.append({
                "text": r.text,
                "font": typeface,
                "size": round(f.size.pt, 1) if f.size else None,
                "color": _color_of(f.color),
                "bold": f.bold,
                "italic": f.italic,
                "align": str(p.alignment).split(" ")[0] if p.alignment else None,
                "line_spacing": p.line_spacing if isinstance(p.line_spacing, float) else None,
            })
    return out


def _shape_type(shape):
    try:
        st = str(shape.shape_type).split(" ")[0]
    except (AttributeError, ValueError):
        st = "?"
    if st in ("AUTO_SHAPE", "AUTO_SHAPE_TYPE"):
        try:
            return str(shape.auto_shape_type).split(" ")[0]
        except (AttributeError, ValueError):
            return st
    return st


def describe(shape, index):
    rec = {
        "index": index,
        "name": shape.name,
        "type": _shape_type(shape),
        "left": _in(shape.left), "top": _in(shape.top),
        "width": _in(shape.width), "height": _in(shape.height),
        "rotation": round(shape.rotation, 1) if getattr(shape, "rotation", 0) else None,
    }
    if shape.shape_type is not None and str(shape.shape_type).startswith("GROUP"):
        rec["children"] = [describe(c, "%s.%d" % (index, i + 1))
                           for i, c in enumerate(shape.shapes)]
        return rec
    rec["fill"] = _fill_of(shape)
    rec["line"] = _line_of(shape)
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        rec["runs"] = _runs_of(shape.text_frame)
        try:
            rec["anchor"] = str(shape.text_frame.vertical_anchor).split(" ")[0] \
                if shape.text_frame.vertical_anchor else None
        except (AttributeError, ValueError):
            rec["anchor"] = None
    if getattr(shape, "has_table", False):
        table = shape.table
        rec["table"] = {
            "rows": len(table.rows), "cols": len(table.columns),
            "col_widths": [_in(c.width) for c in table.columns],
            "row_heights": [_in(r.height) for r in table.rows],
            "first_row": table.first_row, "banding": table.horz_banding,
            "cells": [[{"text": table.cell(r, c).text,
                        "fill": _fill_of(table.cell(r, c))}
                       for c in range(len(table.columns))]
                      for r in range(len(table.rows))],
        }
    if shape.shape_type is not None and str(shape.shape_type).startswith("PICTURE"):
        try:
            img = shape.image
            rec["image"] = {"px": list(img.size), "ext": img.ext,
                            "dpi": list(img.dpi), "bytes": len(img.blob)}
        except (AttributeError, ValueError):
            rec["image"] = {"px": None}
    return rec


def _print_runs(runs, indent):
    for r in runs:
        bits = []
        if r["font"]:
            bits.append(r["font"])
        if r["size"]:
            bits.append("%.1fpt" % r["size"])
        if r["color"]:
            bits.append(r["color"])
        if r["bold"]:
            bits.append("bold")
        if r["italic"]:
            bits.append("italic")
        if r["align"] and r["align"] != "LEFT":
            bits.append(r["align"].lower())
        if r["line_spacing"]:
            bits.append("ls%.2f" % r["line_spacing"])
        text = r["text"].replace("\n", " / ").replace("\x0b", " / ")
        if len(text) > 110:
            text = text[:107] + "..."
        print("%s%-72s  [%s]" % (indent, '"%s"' % text, " ".join(bits)))


def _print_shape(rec, text_only, indent=""):
    if text_only:
        for r in rec.get("runs", []):
            print("%s%s" % (indent, r["text"].replace("\x0b", " / ")))
        if rec.get("table"):
            for row in rec["table"]["cells"]:
                print("%s| %s |" % (indent, " | ".join(c["text"] for c in row)))
        for child in rec.get("children", []):
            _print_shape(child, text_only, indent)
        return
    geo = "L%s T%s W%s H%s" % (_fmt_in(rec["left"]), _fmt_in(rec["top"]),
                              _fmt_in(rec["width"]), _fmt_in(rec["height"]))
    extra = []
    if rec.get("fill") and rec["fill"] != "inherit":
        extra.append("fill=%s" % rec["fill"])
    if rec.get("line") and rec["line"] != "none":
        extra.append("line=%s" % rec["line"])
    if rec.get("rotation"):
        extra.append("rot=%s" % rec["rotation"])
    if rec.get("anchor") and rec["anchor"] != "TOP":
        extra.append("anchor=%s" % rec["anchor"].lower())
    if rec.get("image"):
        extra.append("png=%sx%s" % tuple(rec["image"]["px"] or ("?", "?")))
    print("%s[%s] %-22s %-26s %s  %s" % (indent, rec["index"], rec["type"],
                                         '"%s"' % rec["name"][:24], geo,
                                         "  ".join(extra)))
    _print_runs(rec.get("runs", []), indent + "        ")
    if rec.get("table"):
        t = rec["table"]
        print("%s        TABLE %dx%d  cols=%s  first_row=%s banding=%s"
              % (indent, t["rows"], t["cols"],
                 [("%.2f" % c) if c else "?" for c in t["col_widths"]],
                 t["first_row"], t["banding"]))
        for ri, row in enumerate(t["cells"]):
            cells = " | ".join("%s%s" % (c["text"][:28],
                                         ("(%s)" % c["fill"]) if c["fill"] and
                                         c["fill"] not in ("none", "inherit") else "")
                               for c in row)
            print("%s          r%d: %s" % (indent, ri, cells))
    for child in rec.get("children", []):
        _print_shape(child, text_only, indent + "    ")


def inspect(path, slide_no=None):
    prs = Presentation(str(path))
    out = {
        "file": str(path),
        "canvas_emu": [prs.slide_width, prs.slide_height],
        "canvas_in": [_in(prs.slide_width), _in(prs.slide_height)],
        "slide_count": len(prs.slides),
        "slides": [],
    }
    for i, slide in enumerate(prs.slides, start=1):
        if slide_no and i != slide_no:
            continue
        out["slides"].append({
            "number": i,
            "layout": slide.slide_layout.name,
            "shape_count": len(slide.shapes),
            "shapes": [describe(s, str(j + 1)) for j, s in enumerate(slide.shapes)],
        })
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("pptx")
    ap.add_argument("--slide", type=int, help="only this 1-based slide number")
    ap.add_argument("--text-only", action="store_true", help="print copy only")
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args(argv)

    path = Path(args.pptx).expanduser()
    if not path.exists():
        print("no such file: %s" % path, file=sys.stderr)
        return 2
    data = inspect(path, args.slide)
    if args.json:
        print(json.dumps(data, indent=2))
        return 0
    print("%s" % data["file"])
    print("canvas %d x %d EMU  =  %.3f x %.3f in   |   %d slides"
          % (data["canvas_emu"][0], data["canvas_emu"][1],
             data["canvas_in"][0], data["canvas_in"][1], data["slide_count"]))
    for s in data["slides"]:
        print("")
        print("=" * 78)
        print("SLIDE %d   layout=%s   %d top-level shapes"
              % (s["number"], s["layout"], s["shape_count"]))
        print("=" * 78)
        for rec in s["shapes"]:
            _print_shape(rec, args.text_only)
    return 0


if __name__ == "__main__":
    sys.exit(main())
