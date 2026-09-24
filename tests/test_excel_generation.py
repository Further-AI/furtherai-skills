"""Exercise the Excel skill helpers with synthetic records and workbooks."""

import datetime as dt
import shutil
import subprocess
from pathlib import Path
from unittest.mock import patch

import pytest
import recalc
from check_workbook import check_workbook
from openpyxl import load_workbook
from reconcile import reconcile
from xlsx_kit import (
    CellValue,
    Column,
    ColumnKind,
    MatrixGroup,
    MatrixRow,
    Section,
    SummaryLine,
    add_matrix_sheet,
    add_table_sheet,
    coerce,
    count_rows,
    new_workbook,
    write_summary,
)


@pytest.mark.parametrize(
    ("raw", "kind", "expected"),
    [
        ("001234", "id", "001234"),
        ("$1,234.50", "currency", 1234.5),
        ("(300)", "currency", -300),
        ("94%", "percent", 0.94),
        ("2024", "year", 2024),
        ("09/23/2026", "date", dt.date(2026, 9, 23)),
        ("N/A", "number", None),
        ("<empty>", "text", None),
        ("unreadable", "currency", "unreadable"),
    ],
)
def test_coerce_preserves_identifiers_and_converts_values(
    raw: object, kind: ColumnKind, expected: CellValue
) -> None:
    assert coerce(raw, kind) == expected


def test_table_round_trip_preserves_types_and_formula(tmp_path: Path) -> None:
    workbook = new_workbook()
    sheet = add_table_sheet(
        workbook,
        "Records",
        [Column("Unit #", "id"), Column("Premium", "currency")],
        [["0001", "$1,234.50"], ["0002", None]],
    )
    write_summary(
        workbook["Summary"],
        title="Record summary",
        sections=[
            Section("Results", [SummaryLine("Records", count_rows(sheet.title))])
        ],
    )
    path = tmp_path / "records.xlsx"
    workbook.save(path)
    saved = load_workbook(path)
    assert saved["Records"]["A2"].value == "0001"
    assert saved["Records"]["A2"].data_type == "s"
    assert saved["Records"]["B2"].value == pytest.approx(1234.5)
    assert saved["Records"]["B3"].value is None
    assert saved["Records"].auto_filter.ref == "A1:B3"
    assert saved["Records"].freeze_panes == "A2"
    assert saved["Summary"]["B4"].data_type == "f"
    issues = check_workbook(path)
    assert not [issue for issue in issues if issue.severity == "ERROR"]
    assert any(issue.rule == "uncached-formulas" for issue in issues)
    saved.close()


def test_checker_flags_numeric_text_and_leaked_tokens(tmp_path: Path) -> None:
    workbook = new_workbook()
    sheet = workbook.create_sheet("Records")
    sheet.append(["Premium", "Notes"])
    sheet.append(["$1,200", "<empty>"])
    path = tmp_path / "invalid.xlsx"
    workbook.save(path)
    errors = {issue.rule for issue in check_workbook(path) if issue.severity == "ERROR"}
    assert {"numbers-as-text", "leaked-token"} <= errors


def test_matrix_allows_coverage_text_beside_amounts(tmp_path: Path) -> None:
    workbook = new_workbook()
    add_matrix_sheet(
        workbook,
        "Comparison",
        row_header="Coverage",
        columns=["Expiring", "Renewal"],
        baseline="Expiring",
        groups=[MatrixGroup("Limits", [MatrixRow("Property", [100000, "Excluded"])])],
    )
    path = tmp_path / "comparison.xlsx"
    workbook.save(path)
    issues = check_workbook(path)
    assert not [issue for issue in issues if issue.severity == "ERROR"]
    assert not [
        issue for issue in issues if issue.rule in {"mixed-types", "autofilter"}
    ]


def test_reconcile_accounts_for_duplicates_missing_keys_and_conflicts() -> None:
    left = [
        {"key": "A-100", "unit": "1"},
        {"key": "A100", "unit": "2"},
        {"key": None, "unit": "3"},
        {"key": "LEFT", "unit": "4"},
        {"key": "ONLY", "unit": "5"},
    ]
    right = [
        {"key": "a100", "unit": "1"},
        {"key": "B200", "unit": "3"},
        {"key": "RIGHT", "unit": "4"},
    ]
    result = reconcile(
        left,
        right,
        primary=lambda row: row["key"],
        secondary=lambda row: (row["unit"],),
    )
    assert result.accounted_for(len(left), len(right))
    assert [pair.left["unit"] for pair in result.matched] == ["1", "3"]
    assert [pair.left["unit"] for pair in result.review] == ["4"]
    assert [row["unit"] for row in result.left_only] == ["2", "5"]
    assert result.left_duplicates == {"A100": left[:2]}
    assert result.right_only == []


def test_reconcile_near_keys_require_review() -> None:
    result = reconcile(
        [{"key": "ABC123456"}], [{"key": "ABC123457"}], primary=lambda row: row["key"]
    )
    assert result.matched == []
    assert len(result.review) == 1
    assert result.accounted_for(1, 1)


def test_recalculate_rejects_changed_values_without_overwriting(tmp_path: Path) -> None:
    original = tmp_path / "original.xlsx"
    workbook = new_workbook()
    workbook["Summary"]["A1"] = "Preserve this value"
    workbook.save(original)
    before = original.read_bytes()
    workbook["Summary"]["A1"] = "Changed value"
    changed = tmp_path / "changed.xlsx"
    workbook.save(changed)

    def convert(
        command: list[str], **kwargs: object
    ) -> subprocess.CompletedProcess[bytes]:
        """Simulate a converter that changes a literal cell."""
        destination = Path(command[command.index("--outdir") + 1])
        destination.mkdir()
        shutil.copyfile(changed, destination / original.name)
        return subprocess.CompletedProcess(command, 0)

    with (
        patch.object(recalc.subprocess, "run", side_effect=convert),
        pytest.raises(RuntimeError, match="Value changed"),
    ):
        recalc.recalculate(original, soffice="soffice", timeout_seconds=5)
    assert original.read_bytes() == before


def test_recalc_missing_libreoffice_reports_unavailable(tmp_path: Path) -> None:
    with patch.object(recalc, "find_soffice", return_value=None):
        assert recalc.main([str(tmp_path / "workbook.xlsx")]) == 2
